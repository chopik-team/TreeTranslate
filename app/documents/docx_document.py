"""Modify only translatable Word text; preserve all other ZIP members verbatim.

Runs are allocation slots, never individual translation requests. Field results,
URLs, tabs/breaks and opaque XML are protected boundaries. Inline allocation uses
word boundaries in proportion to the original spans; it is not semantic alignment.
"""
from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZipFile, BadZipFile

from docx import Document
from lxml import etree

from app.documents.errors import InvalidDocumentError, SourceChangedError

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
URL = re.compile(r"(?:https?://|www\.|mailto:)[^\s<>]+|[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}")
MAX_SEGMENT_CHARS = 1200
MAX_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
OPAQUE = {W + x for x in ("del", "fldSimple", "drawing", "object", "pict", "txbxContent", "sdt")}


@dataclass
class Segment:
    text: str
    translated: str | None = None


class TextGroup:
    def __init__(self, nodes):
        self.nodes = tuple(nodes)
        self.original = "".join(n.text or "" for n in nodes)
        self.pieces = []
        cursor = 0
        for match in URL.finditer(self.original):
            self._text(self.original[cursor:match.start()])
            self.pieces.append(match.group())
            cursor = match.end()
        self._text(self.original[cursor:])

    def _text(self, text):
        while text:
            stop = min(len(text), MAX_SEGMENT_CHARS)
            if stop < len(text):
                boundaries = [m.end() for m in re.finditer(r"[.!?。！？](?:\s|$)|\s+", text[:stop])]
                stop = next((p for p in reversed(boundaries) if p >= stop // 2), stop)
            piece, text = text[:stop], text[stop:]
            self.pieces.append(Segment(piece) if any(c.isalpha() for c in piece) else piece)

    @property
    def segments(self):
        return [p for p in self.pieces if isinstance(p, Segment)]

    def apply(self):
        text = "".join(p if isinstance(p, str) else p.translated if p.translated is not None else p.text for p in self.pieces)
        if text == self.original:
            return
        lengths = [len(n.text or "") for n in self.nodes]
        boundaries = [0] + [m.end() for m in re.finditer(r"\s+", text)] + [len(text)]
        # CJK has useful boundaries between ideographs even without spaces.
        if not any(c.isspace() for c in text):
            boundaries = list(range(len(text) + 1))
        offset, cumulative = 0, 0
        for index, (node, length) in enumerate(zip(self.nodes, lengths)):
            cumulative += length
            expected = round(len(text) * cumulative / max(1, len(self.original)))
            end = len(text) if index == len(self.nodes) - 1 else min((b for b in boundaries if b >= offset), key=lambda b: abs(b - expected))
            node.text = text[offset:end]
            node.set(XML_SPACE, "preserve")
            offset = end


class DocxDocument:
    def __init__(self, path: Path):
        self.path = Path(path)
        try:
            if self.path.stat().st_size > MAX_UNCOMPRESSED_BYTES:
                raise InvalidDocumentError()
            self.source_bytes = self.path.read_bytes()
            self.source_hash = sha256(self.source_bytes).hexdigest()
            self.roots = {}
            self.groups = []
            with ZipFile(BytesIO(self.source_bytes)) as archive:
                infos = archive.infolist()
                if len(infos) > 20000 or len({i.filename for i in infos}) != len(infos) or sum(i.file_size for i in infos) > MAX_UNCOMPRESSED_BYTES:
                    raise InvalidDocumentError()
                if archive.testzip() is not None:
                    raise InvalidDocumentError()
                document = Document(BytesIO(self.source_bytes))
                self.structure = (len(document.sections), len(document.tables))
                for info in infos:
                    if info.filename == "word/document.xml" or re.fullmatch(r"word/(?:header|footer)\d+\.xml", info.filename):
                        parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
                        root = etree.fromstring(archive.read(info), parser)
                        self.roots[info.filename] = root
                        self._extract(root)
            self.segments = [segment for group in self.groups for segment in group.segments]
        except (OSError, ValueError, KeyError, BadZipFile, etree.LxmlError):
            raise InvalidDocumentError() from None

    def _extract(self, root):
        field_depth = 0
        for paragraph in root.iter(W + "p"):
            if any(parent.tag in OPAQUE for parent in paragraph.iterancestors()):
                continue
            nodes = []
            def flush():
                if nodes:
                    self.groups.append(TextGroup(nodes))
                    nodes.clear()
            for node in paragraph.iter():
                if node is paragraph:
                    continue
                ancestors = list(node.iterancestors())
                if next((p for p in ancestors if p.tag == W + "p"), None) is not paragraph:
                    continue
                if node.tag in OPAQUE:
                    flush()
                if any(p.tag in OPAQUE for p in ancestors):
                    continue
                if node.tag == W + "fldChar":
                    flush()
                    kind = node.get(W + "fldCharType")
                    field_depth = field_depth + 1 if kind == "begin" else max(0, field_depth - 1) if kind == "end" else field_depth
                elif node.tag in {W + "br", W + "tab", W + "cr", W + "instrText"}:
                    flush()
                elif node.tag == W + "t" and not field_depth:
                    nodes.append(node)
            flush()

    def sample(self):
        if not self.segments:
            return ""
        count = min(24, len(self.segments))
        indices = {round(i * (len(self.segments) - 1) / max(1, count - 1)) for i in range(count)}
        return "\n".join(self.segments[i].text[:320] for i in sorted(indices))

    def assert_source_unchanged(self):
        try:
            if sha256(self.path.read_bytes()).hexdigest() != self.source_hash:
                raise SourceChangedError()
        except OSError:
            raise SourceChangedError() from None

    def write(self, output: Path):
        for group in self.groups:
            group.apply()
        with ZipFile(BytesIO(self.source_bytes)) as original, ZipFile(output, "w") as result:
            result.comment = original.comment
            for info in original.infolist():
                data = etree.tostring(self.roots[info.filename], encoding="UTF-8", xml_declaration=True, standalone=True) if info.filename in self.roots else original.read(info)
                result.writestr(info, data)


def validate_docx(path: Path, expected_structure=None):
    try:
        if not path.is_file() or path.stat().st_size == 0:
            raise InvalidDocumentError()
        with ZipFile(path) as archive:
            if archive.testzip() is not None:
                raise InvalidDocumentError()
        document = Document(path)
        structure = len(document.sections), len(document.tables)
        if expected_structure is not None and structure != expected_structure:
            raise InvalidDocumentError()
        # Materialize cell/section proxies so malformed tables/sections fail validation.
        for table in document.tables:
            for row in table.rows:
                tuple(row.cells)
        tuple(document.sections)
        return structure
    except (OSError, ValueError, KeyError, BadZipFile, etree.LxmlError):
        raise InvalidDocumentError() from None
