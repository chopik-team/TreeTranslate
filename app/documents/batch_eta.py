"""Approximate workload budgets, refined by progress; no OCR/model pretranslation."""
import logging
from threading import RLock

logger = logging.getLogger('treetranslate.documents.job')


class BatchEta:
    def __init__(self, rates, control):
        self.rates, self.control = rates, control
        self.files = []
        self.catalog = {}
        self.workloads = []
        self.segment_counts = {}
        self.phase = None
        self.phase_started = 0
        self.available = True
        self.translation_rate = rates['segment_seconds']
        self._lock = RLock()
        self._signature = None
        self._archive = False
        self._archive_fraction = 0.
        self._archive_anchor = None

    @staticmethod
    def key(item):
        return str(item.path), item.size

    def _inspect(self, item, config):
        from app.documents.pdf_document import PDF_LOCK
        import pypdfium2 as pdfium
        import pypdfium2.raw as raw
        from app.ocr.config import configuration
        minimum = configuration()['native']['minimum_region_points']
        pages, chars, regions = 1, 0, 0
        if item.path.suffix.lower() == '.pdf':
            with PDF_LOCK, pdfium.PdfDocument(item.path) as pdf:
                pages = len(pdf)
                # A bounded sample for large PDFs; no render, OCR or semantic pipeline.
                indices = sorted({0, max(0,pages//2), max(0,pages-1)})
                for index in indices:
                    self.control.checkpoint()
                    page = pdf[index]
                    try:
                        text = page.get_textpage()
                        try:native = len(text.get_text_range().strip())
                        finally:text.close()
                        chars += native
                        if config.ocr_enabled:
                            bounds = page.get_bbox()
                            area = max(1, (bounds[2]-bounds[0])*(bounds[3]-bounds[1]))
                            for obj in page.get_objects():
                                if obj.type != raw.FPDF_PAGEOBJ_IMAGE:continue
                                x,y,right,top = obj.get_bounds()
                                if right-x >= minimum and top-y >= minimum and not (native >= 20 and (right-x)*(top-y)/area >= .65):
                                    regions += 1
                    finally:page.close()
                scale = pages/max(1,len(indices))
                chars *= scale;regions *= scale
        else:
            from docx import Document
            doc = Document(item.path)
            chars = sum(len(p.text) for p in doc.paragraphs)
            chars += sum(len(cell.text) for table in doc.tables for row in table.rows for cell in row.cells)
        kind = 'image_only' if regions and chars<20*pages else 'mixed' if regions else 'native'
        return dict(pages=pages,chars=chars,regions=regions,kind=kind)

    def _budget(self, item, workload):
        rates = self.rates
        if workload is None:
            # Archive inventory already has member sizes. Never open/extract all paths for ETA.
            cost = rates.get('archive_seconds_per_size_unit', rates['segment_seconds']*rates.get('segments_per_page',10))
            total = max(1,cost*max(.25,(item.size/rates.get('size_unit_bytes',32768))**.5))
            return dict(extract=total*.35,translate=total*.53,write=total*.12)
        profile = rates.get('classes',{}).get(workload['kind'],rates)
        segments = max(1,workload['chars']/max(1,profile.get('chars_per_segment',80)) + workload['regions']*profile.get('segments_per_page',10))
        budget = dict(extract=workload['regions']*(profile.get('ocr_seconds') or 0),
                      translate=segments*profile['segment_seconds'],
                      write=max(1,workload['pages']*profile['write_page_seconds']))
        floor = workload['pages']*profile.get('document_page_seconds',0)
        budget['extract'] += max(0,floor-sum(budget.values()))
        return budget

    def prepare(self, files, config, run_id):
        signature = tuple(self.key(f) for f in files),config.ocr_enabled
        with self._lock:
            if signature == self._signature:return
        budgets,workloads = [],[]
        inspected = 0
        for index,item in enumerate(files,1):
            self.control.checkpoint()
            key = (*self.key(item),config.ocr_enabled)
            workload = self.catalog.get(key)
            if key not in self.catalog and not item.archive and inspected<16:
                inspected += 1
                try:workload = self._inspect(item,config)
                except Exception as error:
                    from app.engine.errors import TranslationCancelledError
                    if isinstance(error,TranslationCancelledError):raise
                    logger.info('run=%s workload file=%d inventory=unavailable type=%s',run_id,index,type(error).__name__)
                self.catalog[key] = workload
            if workload and workload['regions'] and self.rates.get('ocr_seconds') is None and not self.rates.get('classes'):
                self.available = False
            budgets.append(self._budget(item,workload));workloads.append(workload)
            if workload:logger.info('run=%s workload file=%d native_chars=%d ocr_regions=%d pages=%d',
                                   run_id,index,workload['chars'],workload['regions'],workload['pages'])
        with self._lock:
            self.files,self.workloads = budgets,workloads
            self.segment_counts = {}
            self._signature=signature;self._archive=any(f.archive for f in files)
            self.phase=None;self._archive_anchor=None;self._archive_fraction=0.
        logger.info('run=%s eta calibration_runs=%s estimated_seconds=%.2f',run_id,self.rates.get('runs',[]),self.remaining())

    def begin(self):
        with self._lock:
            self.phase=(1,'extract') if self.files else None
            self.phase_started=self.control.active_seconds
            if self._archive:self._archive_anchor=(self.remaining(),self.control.active_seconds)

    def extracted(self, index, segments):
        with self._lock:
            self.segment_counts[index] = segments
            self.files[index-1]['extract']=0
            workload=self.workloads[index-1] if self.workloads else None
            profile=self.rates.get('classes',{}).get((workload or {}).get('kind'),self.rates)
            self.files[index-1]['translate']=segments*profile['segment_seconds']
            self.phase=None

    def published(self, index):
        with self._lock:
            self.files[index-1]=dict(extract=0,translate=0,write=0)
            self.phase=None

    def update(self, progress, durations):
        with self._lock:
            if not self.available or not self.files:return
            completed = progress.stage in {'COMPLETED','COMPLETED_WITH_FAILURES'}
            if self._archive:
                # Canonical publication progress may interleave prepared children. Use a watermark.
                fraction=min(.99,max(0,progress.percent/100))
                if fraction>self._archive_fraction and self.control.active_seconds>=5:
                    prior=sum(sum(f.values()) for f in self.files)
                    observed=self.control.active_seconds/max(.01,fraction)
                    weight=min(.8,fraction)
                    remaining=(1-fraction)*(prior*(1-weight)+observed*weight)
                    self._archive_anchor=(remaining,self.control.active_seconds)
                    self._archive_fraction=fraction
            elif progress.file_index and progress.file_index<=len(self.files):
                phase = 'extract' if progress.stage in {'EXTRACTING','RENDERING','OCR','LAYOUT_ANALYSIS'} else 'translate' if progress.stage=='TRANSLATING' else 'write'
                key=(progress.file_index,phase)
                workload=self.workloads[progress.file_index-1] if self.workloads else None
                profile=self.rates.get('classes',{}).get((workload or {}).get('kind'),self.rates)
                prior=profile['segment_seconds']
                if self.phase!=key:
                    self.phase,self.phase_started=key,self.control.active_seconds
                    if phase=='translate':self.translation_rate=prior
                if phase=='translate' and len(durations)>=3:
                    rate=max(prior*.25,prior*.7+sum(durations)/len(durations)*.3)
                    current=self.files[progress.file_index-1]
                    if all(i in self.segment_counts for i in range(1,progress.file_index+1)):
                        # Segment totals are global, but the measured rate belongs
                        # to this document. Keep other complexity classes' priors.
                        remaining=max(0,sum(self.segment_counts[i] for i in range(1,progress.file_index+1))-progress.processed)
                        current['translate']=remaining*rate
                    else:
                        remaining=max(0,progress.total-progress.processed)
                        others=sum(f['translate'] for f in self.files[progress.file_index:])
                        current['translate']=max(0,remaining*rate-others)
                    self.translation_rate=rate
                    self.phase_started=self.control.active_seconds
                elif phase=='write':self.files[progress.file_index-1]['translate']=0
            if completed:
                self.files=[dict(extract=0,translate=0,write=0) for _ in self.files]
                self.phase=None;self._archive_anchor=(0,self.control.active_seconds)
            progress.eta_scope='batch'
            progress.eta_seconds=0 if completed else max(0,round(self.remaining()))

    def remaining(self):
        with self._lock:
            if self._archive_anchor is not None:
                total,started=self._archive_anchor
                return max(0,total-(self.control.active_seconds-started))
            total=sum(sum(f.values()) for f in self.files)
            if self.phase:
                total-=max(0,self.control.active_seconds-self.phase_started)
            return max(0,total)
