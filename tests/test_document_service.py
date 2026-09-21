import os
from types import SimpleNamespace
from hashlib import sha256
from unittest.mock import patch
from pathlib import Path
from time import monotonic, sleep

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
from docx import Document
from PySide6.QtWidgets import QApplication
from PySide6.QtCore import QSettings

from app.documents.job import DocumentConfig
from app.gui.main_window import MainWindow
from app.models.translation_job import JobState
from app.services.hybrid_translation_service import HybridTranslationService
from app.services.settings_service import SettingsService
from test_hybrid_service import SlowEngine


def wait_for(predicate, timeout=3000):
    deadline = monotonic() + timeout / 1000
    while not predicate() and monotonic() < deadline:
        QApplication.processEvents()
        # qWait can retain the GIL on Windows; let the Python document worker run.
        sleep(.005)
    QApplication.processEvents()
    assert predicate()


def make_source(tmp_path):
    path = tmp_path / 'test.docx'
    doc = Document()
    doc.add_paragraph('Private paragraph one')
    doc.add_paragraph('Private paragraph two')
    doc.save(path)
    return path


def make_engine():
    engine = SlowEngine()
    engine.languages = SimpleNamespace(resolve=lambda text, source, target: ('en', 'ru'))
    return engine


def test_file_cancel_retains_session_until_worker_finishes(tmp_path):
    app = QApplication.instance() or QApplication([])
    source = make_source(tmp_path)
    digest = sha256(source.read_bytes()).hexdigest()
    engine = make_engine()
    service = HybridTranslationService(engine=engine)
    window = MainWindow(translation_service=service)
    errors = []
    service.files.failed.connect(errors.append)
    try:
        window.translation.accept_paths([source])
        wait_for(lambda: service.state == JobState.READY)
        window.translation._start_translation()
        assert engine.entered.wait(2)
        service.cancel()
        assert service.state == JobState.CANCELLING
        assert window.sessions.is_active('file')
        window.translation.translate_text('must not run')
        assert not window.sessions.is_active('text')
        assert len(engine.calls) == 1
        engine.release.set()
        wait_for(lambda: service.state == JobState.CANCELLED)
        assert not window.sessions.is_active('file')
        assert not window.file_page.progress.output_paths
        assert sha256(source.read_bytes()).hexdigest() == digest
    finally:
        engine.release.set()
        window.close()


def test_file_service_completed_output_and_open_actions(tmp_path):
    app = QApplication.instance() or QApplication([])
    source = make_source(tmp_path)
    engine = make_engine()
    engine.release.set()
    service = HybridTranslationService(engine=engine)
    window = MainWindow(translation_service=service)
    errors = []
    service.files.failed.connect(errors.append)
    try:
        window.translation.accept_paths([source])
        wait_for(lambda: service.state == JobState.READY)
        service.files.config = DocumentConfig(source='en', output=tmp_path / 'result')
        service.start()
        wait_for(lambda: service.state in {JobState.COMPLETED, JobState.ERROR})
        assert service.state == JobState.COMPLETED, errors
        output, = window.file_page.progress.output_paths
        assert output.is_file()
        assert window.file_page.progress.bar.value() == 100
        with patch('app.controllers.translation_ui_controller.QDesktopServices.openUrl', return_value=True) as open_url:
            window.file_page.progress.open_file.click()
            assert Path(open_url.call_args.args[0].toLocalFile()) == output.resolve()
        with patch('app.controllers.translation_ui_controller.QProcess.startDetached', return_value=True) as explorer:
            window.file_page.progress.show_output.click()
            assert explorer.call_args.args[0] == 'explorer.exe'
            assert '/select,' in explorer.call_args.args[1]
    finally:
        window.close()


def test_pdf_scan_is_honest_and_cannot_start(tmp_path):
    app = QApplication.instance() or QApplication([])
    pdf = tmp_path / 'manual.pdf'
    pdf.write_bytes(b'%PDF')
    engine = make_engine()
    engine.release.set()
    window = MainWindow(translation_service=HybridTranslationService(engine=engine))
    try:
        window.translation.accept_paths([pdf])
        wait_for(lambda: window.translation_service.state == JobState.READY)
        assert not window.file_page.start_button.isEnabled()
        assert 'Формат пока не поддерживается' in window.file_page.file_status.text()
        assert not engine.calls
        assert window.file_page.progress.bar.value() == 0
    finally:
        window.close()


def test_completed_job_opens_output_folder_only_when_enabled(tmp_path):
    app = QApplication.instance() or QApplication([])
    window = MainWindow(translation_service=HybridTranslationService(engine=make_engine()))
    output = tmp_path / 'result' / 'translated.docx'
    output.parent.mkdir()
    output.write_bytes(b'result')
    try:
        window.file_page.progress.set_output_paths([output])
        with patch.object(window.translation.settings, 'value', return_value=False), \
             patch('app.controllers.translation_ui_controller.QDesktopServices.openUrl', return_value=True) as open_url:
            window.translation._state_changed(JobState.COMPLETED)
            open_url.assert_not_called()

        with patch.object(window.translation.settings, 'value', return_value=True), \
             patch('app.controllers.translation_ui_controller.QDesktopServices.openUrl', return_value=True) as open_url:
            window.translation._state_changed(JobState.COMPLETED)
        open_url.assert_called_once()
        assert Path(open_url.call_args.args[0].toLocalFile()) == output.parent.resolve()
    finally:
        window.close()


def test_unfinished_job_is_restored_as_ready_without_auto_start(tmp_path):
    app = QApplication.instance() or QApplication([])
    source = make_source(tmp_path)
    settings = SettingsService(QSettings(str(tmp_path / 'settings.ini'), QSettings.Format.IniFormat))
    settings.save_value('general/restore_job', True)
    settings.save_unfinished_job([source])
    engine = make_engine()
    engine.release.set()
    service = HybridTranslationService(engine=engine)
    window = MainWindow(translation_service=service)
    window.translation.settings = settings
    window.translation.preferences.settings = settings
    try:
        window.show()
        wait_for(lambda: service.state == JobState.READY)
        assert window.file_page.start_button.isEnabled()
        assert window.file_page.file_tree.selected_paths() == [source.resolve()]
        assert 'Задача восстановлена' in window.file_page.file_status.text()
        assert not engine.calls
    finally:
        window.close()


def test_finished_or_cancelled_job_is_not_restored(tmp_path):
    app = QApplication.instance() or QApplication([])
    source = make_source(tmp_path)
    settings = SettingsService(QSettings(str(tmp_path / 'settings.ini'), QSettings.Format.IniFormat))
    settings.save_value('general/restore_job', True)
    settings.save_unfinished_job([source])
    window = MainWindow(translation_service=HybridTranslationService(engine=make_engine()))
    window.translation.settings = settings
    try:
        window.translation._state_changed(JobState.CANCELLED)
        assert settings.unfinished_job_paths() == ()
    finally:
        window.close()


def test_completed_job_stays_cleared_until_restarted(tmp_path):
    app = QApplication.instance() or QApplication([])
    source = make_source(tmp_path)
    settings = SettingsService(QSettings(str(tmp_path / 'settings.ini'), QSettings.Format.IniFormat))
    settings.save_value('general/restore_job', True)
    engine = make_engine()
    engine.release.set()
    service = HybridTranslationService(engine=engine)
    window = MainWindow(translation_service=service)
    window.translation.settings = settings
    try:
        window.translation.accept_paths([source])
        wait_for(lambda: service.state == JobState.READY)
        window.translation._start_translation()
        wait_for(lambda: service.state == JobState.COMPLETED)
        window.translation.update_restore_preference()
        assert settings.unfinished_job_paths() == ()
        window.translation._start_translation()
        assert settings.unfinished_job_paths() == (source.resolve(),)
        wait_for(lambda: service.state == JobState.COMPLETED)
    finally:
        window.close()


def test_shutdown_preserves_active_job_and_ignores_rejected_paths(tmp_path):
    app = QApplication.instance() or QApplication([])
    source = make_source(tmp_path)
    settings = SettingsService(QSettings(str(tmp_path / 'settings.ini'), QSettings.Format.IniFormat))
    settings.save_value('general/restore_job', True)
    engine = make_engine()
    service = HybridTranslationService(engine=engine)
    window = MainWindow(translation_service=service)
    window.translation.settings = settings
    try:
        window.translation.accept_paths([source])
        wait_for(lambda: service.state == JobState.READY)
        window.translation._start_translation()
        assert engine.entered.wait(2)
        window.translation.accept_paths([tmp_path / 'rejected.docx'])
        assert settings.unfinished_job_paths() == (source.resolve(),)
    finally:
        engine.release.set()
        window.close()
    QApplication.processEvents()
    assert settings.unfinished_job_paths() == (source.resolve(),)
