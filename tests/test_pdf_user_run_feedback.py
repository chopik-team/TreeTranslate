from pathlib import Path
from types import SimpleNamespace
from PySide6.QtCore import Qt
from app.gui.widgets.file_tree import FileTree
from app.gui.widgets.progress_panel import ProgressPanel
from app.models.file_item import FileItem
from app.models.translation_job import JobState, TranslationProgress
from app.ocr.pdf_extractor import HybridPdfExtractor
from app.documents.job import DocumentConfig


def test_ocr_label_does_not_wrap_into_neighbouring_column(tmp_path):
    from app.documents.pdf_document import PdfDocument
    from app.documents.pdf_types import PdfSegment
    from tools.pdf_fixtures import make_pdf
    doc = PdfDocument(make_pdf(tmp_path / 'source.pdf', image=True))
    for segment in doc.segments:
        segment.translated = segment.text
    left = PdfSegment(0, 'left-label', '展开全部', (450, 500, 485, 510), 10, (), 10,
                      origin='ocr', available_bbox=(445, 490, 497, 514), translated='Откройте все')
    right = PdfSegment(0, 'right-label', '系统', (500, 500, 520, 510), 11, (), 10,
                       origin='ocr', available_bbox=(475, 490, 535, 514), translated='Система')
    doc.segments.extend((left, right))
    out = tmp_path / 'result.pdf'
    doc.write(out)
    doc.validate(out)
    assert left.overflow_text == left.translated
    for segment in (left, right):
        assert segment.rendered_bbox[0] >= segment.bbox[0]
        assert segment.rendered_bbox[1] >= segment.bbox[1]
        assert segment.rendered_bbox[3] <= segment.bbox[3]


def test_batch_keeps_published_files_and_retries_only_checked_remaining(tmp_path, qt_application, monkeypatch):
    from time import monotonic, sleep
    from docx import Document
    from app.services.hybrid_translation_service import HybridTranslationService
    from app.gui.main_window import MainWindow
    import app.documents.job as module
    from app.documents.errors import DocumentError
    files = [tmp_path / name for name in ('first.docx', 'second.docx')]
    for source in files:
        doc = Document(); doc.add_paragraph('Source ' + source.stem); doc.save(source)
    calls = []
    engine = SimpleNamespace(shutdown=lambda: None,
        languages=SimpleNamespace(resolve=lambda text, source, target: ('en', 'ru')),
        translate=lambda request, cancel: (calls.append(request.text) or SimpleNamespace(translated_text='Перевод ' + request.text)))
    service = HybridTranslationService(engine=engine)
    window = MainWindow(translation_service=service)
    def wait():
        deadline = monotonic() + 5
        while service.files.busy and monotonic() < deadline:
            qt_application.processEvents(); sleep(.01)
        qt_application.processEvents()
        assert not service.files.busy
    original = module.validate_docx
    validations = []
    def validate(path, structure):
        validations.append(path)
        if len(validations) == 2:
            raise DocumentError('Expected test failure')
        return original(path, structure)
    monkeypatch.setattr(module, 'validate_docx', validate)
    try:
        window.translation.accept_paths(files); wait()
        window.file_page.start_button.click(); wait()
        assert service.state == JobState.ERROR
        completed = tuple(window.file_page.progress.output_paths)
        assert len(completed) == 1 and completed[0].exists()
        failed_percent = window.file_page.progress.bar.value()
        assert 0 < failed_percent < 100
        assert files[0] not in window.file_page.file_tree.selected_paths()
        assert files[1] in window.file_page.file_tree.selected_paths()
        window.file_page.start_button.click(); wait()
        assert service.state == JobState.COMPLETED
        assert len(window.file_page.progress.output_paths) == 2
        assert calls.count('Source first') == 1
        assert all(p.exists() for p in window.file_page.progress.output_paths)
        from app.controllers.translation_ui_controller import QMessageBox
        first_item = window.file_page.file_tree.tree.topLevelItem(0).child(0)
        first_item.setCheckState(0, Qt.CheckState.Checked)
        questions = []
        def decline(*args):
            questions.append(args)
            return QMessageBox.StandardButton.No
        monkeypatch.setattr(QMessageBox, 'question', decline)
        window.file_page.start_button.click(); wait()
        assert questions and calls.count('Source first') == 1
        monkeypatch.setattr(QMessageBox, 'question', lambda *args: QMessageBox.StandardButton.Yes)
        window.file_page.start_button.click(); wait()
        assert calls.count('Source first') == 2
        assert len(window.file_page.progress.output_paths) == 3
    finally:
        window.close()


def test_display_name_is_clean_but_source_path_unchanged(qt_application):
    path = 'C:/files/维修程序.pdf.pdf'
    tree = FileTree()
    tree.populate(FileItem('root', True, children=[FileItem('维修程序.pdf.pdf', path=path)]))
    item = tree.tree.topLevelItem(0).child(0)
    assert item.text(0) == '维修程序.pdf'
    assert item.data(0, Qt.ItemDataRole.UserRole) == path
    panel = ProgressPanel()
    panel.set_progress(TranslationProgress(current_file='维修程序.pdf.pdf · страница 2'))
    assert panel.current.text() == '维修程序.pdf · страница 2'


def test_busy_tree_keeps_hover_but_locks_checks_and_completed_files(qt_application):
    source, output = Path('C:/files/manual.pdf'), Path('C:/files/manual_ru.pdf')
    tree = FileTree()
    tree.populate(FileItem('root', True, children=[FileItem(source.name, path=str(source))]))
    item = tree.tree.topLevelItem(0).child(0)
    tree.set_busy(True)
    assert tree.isEnabled()
    assert not item.flags() & Qt.ItemFlag.ItemIsUserCheckable
    tree.mark_completed(source, output)
    tree.set_busy(False)
    assert item.flags() & Qt.ItemFlag.ItemIsUserCheckable
    assert item.checkState(0) == Qt.CheckState.Unchecked
    assert item.data(0, Qt.ItemDataRole.UserRole + 4) == 100
    assert source not in tree.selected_paths()


def test_extraction_eta_uses_warm_observations_and_labels_scope(qt_application):
    seen = []
    extractor = HybridPdfExtractor(None, DocumentConfig(), progress=lambda *values: seen.append(values))
    extractor._page_regions = 2
    extractor._current_region = 0
    page = SimpleNamespace(pdf=[None] * 4)
    extractor.timings = [dict(render_seconds=1, ocr_seconds=900)]
    extractor.report_progress('OCR', page, 1)
    assert seen[-1][-1] is None
    extractor.timings.extend([dict(render_seconds=1, ocr_seconds=29)] * 2)
    extractor.report_progress('OCR', page, 2)
    assert seen[-1][-1] == 120
    panel = ProgressPanel()
    panel.set_state(JobState.TRANSLATING)
    panel.set_progress(TranslationProgress(stage='OCR', total=0, eta_seconds=120, eta_scope='stage'))
    assert panel.remaining_caption.text() == 'До завершения этапа'
    assert panel.remaining.text() == '00:02:00'
