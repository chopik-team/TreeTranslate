"""Real prepared Argos/M2M100 and Paddle OCR; fixture terms only."""
from pathlib import Path
from hashlib import sha256
from unittest.mock import Mock
import json
import os
import socket
from docx import Document
import pypdfium2 as pdfium
import pytest

from app.engine.factory import create_translation_engine
from app.engine.types import TranslationRequest,DevicePreference,PerformanceProfile
from app.translation_memory.engine import TranslationMemoryEngine
from app.glossary.engine import GlossaryEngine
from app.glossary.constraints import validate_result
from app.documents.control import JobControl
from app.documents.job import DocumentConfig,DocumentJob
from app.documents.scanner import scan_sources
from tools.pdf_fixtures import make_pdf
from tools.ocr_fixtures import rasterize


@pytest.mark.integration
def test_real_backend_placeholders_and_enforcement(tmp_path,monkeypatch):
    def fail(*a,**kw):raise AssertionError('Network forbidden')
    monkeypatch.setattr(socket,'getaddrinfo',fail);monkeypatch.setattr(socket,'create_connection',fail)
    g=GlossaryEngine(tmp_path/'glossary.db')
    tm=TranslationMemoryEngine(tmp_path/'tm.db');engine=create_translation_engine(memory=tm,glossary=g)
    rows=[]
    cases=[('argos','en','ru','Replace the coolant and check the connector.', [('coolant','антифриз'),('connector','разъём')]),
           ('argos','ru','en','Замените антифриз и проверьте разъём.', [('антифриз','coolant'),('разъём','connector')]),
           ('m2m100','zh','ru','更换冷却液并检查连接器。',[('冷却液','антифриз'),('连接器','разъём')]),
           ('m2m100','en','ru','Replace the coolant and check the connector.',[('coolant','антифриз'),('connector','разъём')])]
    try:
        for backend,src,tgt,text,terms in cases:
            for source,target in terms:g.remember_term(source,target,src,tgt)
            request=TranslationRequest(text,src,tgt,DevicePreference.AUTO,PerformanceProfile.BALANCED)
            original=engine.translate(request,backend_only=backend)
            router=Mock()
            router.translate.side_effect=lambda req,cancelled:engine.router.translate(req,cancelled,backend_only=backend)
            from threading import Event
            result=g.translate(request,router,Event())
            assert result.constraint_status=='enforced'
            assert all(target in result.translated_text for _,target in terms)
            assert 'ZXQ' not in result.translated_text
            rows.append(dict(backend=backend,source_language=src,target_language=tgt,source=text,
                baseline=original.translated_text,glossary=result.translated_text,device=result.device,status=result.constraint_status,
                model_ids=result.model_ids))
        out=Path(os.environ.get('TREETRANSLATE_QA_DIR','docs/qa'))/'aw071/real-backends.json';out.parent.mkdir(parents=True,exist_ok=True)
        out.write_text(json.dumps(rows,ensure_ascii=False,indent=2),'utf-8')
    finally:engine.shutdown()


@pytest.mark.integration
def test_glossary_docx_pdf_ocr(tmp_path):
    g=GlossaryEngine(tmp_path/'glossary.db');g.remember_term('coolant','антифриз','en','ru',case_sensitive=False)
    tm=TranslationMemoryEngine(tmp_path/'tm.db');engine=create_translation_engine(memory=tm,glossary=g)
    source='Replace the coolant before starting.'
    docx=tmp_path/'manual.docx';document=Document();document.add_paragraph(source);document.save(docx)
    pdf=make_pdf(tmp_path/'native.pdf',source,lines=1,background=False)
    scan=rasterize(pdf,tmp_path/'scan.pdf')
    paths=[docx,pdf,scan];digests={p:sha256(p.read_bytes()).hexdigest() for p in paths}
    seen=[];original=engine.translate
    def record(request,cancelled):
        result=original(request,cancelled);seen.append(result);return result
    try:
        control=JobControl()
        outputs=DocumentJob(scan_sources(paths,control).files,
            DocumentConfig(source='en',target='ru',device=DevicePreference.CPU,profile=PerformanceProfile.FAST,output=tmp_path/'out'),
            control,record,engine.languages.resolve).run()
        assert len(outputs)==3
        for output in outputs:
            if output.suffix=='.docx':text='\n'.join(p.text for p in Document(output).paragraphs)
            else:
                with pdfium.PdfDocument(output) as doc:
                    page=doc[0];tp=page.get_textpage();text=tp.get_text_range();tp.close();page.close()
            assert 'антифриз' in text and 'ZXQ' not in text
        assert sum(r.constraint_status=='enforced' for r in seen)>=3
        assert all(sha256(p.read_bytes()).hexdigest()==digests[p] for p in paths)
        out=Path(os.environ.get('TREETRANSLATE_QA_DIR','docs/qa'))/'aw071/documents.json';out.parent.mkdir(parents=True,exist_ok=True)
        out.write_text(json.dumps(dict(formats=['DOCX','native PDF','OCR PDF'],count=3,glossary_enforced=sum(r.constraint_status=='enforced' for r in seen),
            immutable_originals=True,real_ocr=True,real_translation=True,tm_units=tm.stats()['total_units']),indent=2),'utf-8')
    finally:engine.shutdown()


def test_standalone_pdf_label_bypasses_title_preservation(tmp_path):
    from app.translation_memory.knowledge import TranslationKnowledgeEngine
    from app.engine.router.translation_router import TranslationRouter
    g=GlossaryEngine(tmp_path/'g.db');g.remember_term('Coolant','Антифриз','en','ru')
    tm=TranslationMemoryEngine(tmp_path/'tm.db');tm.remember_translation('User guide','Руководство','en','ru')
    router=TranslationRouter({});router.translate=Mock(side_effect=AssertionError('Direct label must not call model'))
    engine=TranslationKnowledgeEngine(router,tm,g)
    source=make_pdf(tmp_path/'label.pdf','Coolant',lines=1,background=False)
    try:
        control=JobControl()
        output,=DocumentJob(scan_sources([source],control).files,DocumentConfig(source='en',target='ru',output=tmp_path/'out'),
                           control,engine.translate,engine.languages.resolve).run()
        with pdfium.PdfDocument(output) as document:
            page=document[0];tp=page.get_textpage();text=tp.get_text_range();tp.close();page.close()
        assert 'Антифриз' in text
        router.translate.assert_not_called()
    finally:engine.shutdown()
