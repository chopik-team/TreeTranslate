import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from docx import Document
from PySide6.QtCore import QSettings
from PySide6.QtWidgets import QApplication, QLabel, QComboBox

from app.localization import LOCALES, ROOT, localization, saved_locale
from app.gui.main_window import MainWindow
from app.gui.dialogs.settings_dialog import SettingsDialog
from app.services.settings_service import SettingsService
from app.services.mock_translation_service import MockTranslationService
from app.glossary.engine import GlossaryEngine
from app.glossary.bundled import bundled_paths
from app.glossary.domain_detection import DomainDetector
from app.translation_memory.engine import TranslationMemoryEngine
from app.translation_memory.knowledge import TranslationKnowledgeEngine
from app.engine.types import TranslationRequest, TranslationResult

AUTO = 'охлаждающая жидкость и коленчатые валы'
METAL = 'азотирование и цементация стали'


@pytest.fixture
def glossary(tmp_path):
    return GlossaryEngine(tmp_path/'terms.db', builtin_paths=bundled_paths())


def test_actual_pack_evidence_conservative_and_content_free(glossary, monkeypatch, caplog):
    import socket
    monkeypatch.setattr(socket,'create_connection',lambda *a,**k:pytest.fail('network'))
    detector=DomainDetector(glossary)
    assert detector.detect(AUTO,'ru','zh').domain=='automotive'
    assert detector.detect(METAL,'ru','zh').domain=='metallurgy'
    assert detector.detect(AUTO+', '+METAL,'ru','zh').domain=='general'
    assert detector.detect('Система, машина, двигатель, сталь. Добрый день.','ru','zh').domain=='general'
    assert detector.detect('охлаждающая жидкость '*30,'ru','zh').domain=='general'
    assert AUTO not in caplog.text and METAL not in caplog.text
    assert not hasattr(detector,'text')


def test_untrusted_evidence_is_ignored(tmp_path):
    g=GlossaryEngine(tmp_path/'terms.db')
    for term in ('special transmission','special gearbox'):
        g.repository.insert_many([dict(source_term=term,target_term='термин',source_language='en',
            target_language='ru',domain='automotive',status='AUTO')])
    assert DomainDetector(g).detect('special transmission special gearbox','en','ru').domain=='general'


def test_auto_domain_preserves_confirmed_general_tm(glossary,tmp_path):
    memory=TranslationMemoryEngine(tmp_path/'tm.db')
    memory.remember_translation(AUTO,'已确认','ru','zh')
    router=SimpleNamespace(policy=SimpleNamespace(max_text_chars=50000),
        languages=SimpleNamespace(resolve=lambda *a:('ru','zh')),translate=Mock())
    result=TranslationKnowledgeEngine(router,memory,glossary).translate(TranslationRequest(AUTO,'ru','zh',domain='auto'))
    assert result.translated_text=='已确认' and result.backend=='translation_memory'
    assert result.domain=='automotive'
    router.translate.assert_not_called()


def test_batch_detects_once_per_document_after_extraction(glossary,tmp_path):
    from app.documents.job import DocumentJob,DocumentConfig
    from app.documents.scanner import scan_sources
    from app.documents.control import JobControl
    from hashlib import sha256
    paths=[]
    for i,text in enumerate((AUTO,METAL)):
        p=tmp_path/f'{i}.docx';doc=Document();doc.add_paragraph(text);doc.add_paragraph('Ещё один абзац');doc.save(p);paths.append(p)
    hashes=[sha256(p.read_bytes()).hexdigest() for p in paths]
    requests=[]
    class Engine:
        detect_domain=Mock(side_effect=DomainDetector(glossary).detect)
        def translate(self,r,c):
            requests.append(r)
            return SimpleNamespace(translated_text='翻译')
    engine=Engine();control=JobControl()
    job=DocumentJob(scan_sources(paths,control).files,DocumentConfig(source='ru',target='zh',domain='auto',output=tmp_path/'out'),
        control,engine.translate,lambda *a:('ru','zh'))
    job.run()
    assert engine.detect_domain.call_count==2
    assert [r.domain for r in requests]==['automotive']*2+['metallurgy']*2
    assert [sha256(p.read_bytes()).hexdigest() for p in paths]==hashes


def test_document_sample_keeps_late_evidence():
    from app.glossary.domain_detection import document_sample
    paragraphs=[SimpleNamespace(text='filler') for _ in range(80)]
    paragraphs[-1]=SimpleNamespace(text=METAL)
    assert METAL in document_sample(paragraphs)
    assert len(document_sample([SimpleNamespace(text='x'*50000)]))<=8000


@pytest.mark.parametrize('locale',LOCALES)
def test_locale_live_switch_persistence_and_state(locale,tmp_path,monkeypatch):
    app=QApplication.instance() or QApplication([])
    path=str(tmp_path/'settings.ini')
    settings=SettingsService(QSettings(path,QSettings.Format.IniFormat))
    monkeypatch.setattr('app.gui.main_window.SettingsService',lambda:settings)
    window=MainWindow(translation_service=MockTranslationService())
    dialog=SettingsDialog(window,settings,window.preferences)
    try:
        window.preferences.set_languages('Китайский','Русский')
        window.preferences.set_device('CPU')
        window.text_page.source.editor.setPlainText('User content 中文: Настройки')
        window.text_page._translate_timer.stop()
        settings.save_value('general/output_path','C:/Private/Настройки')
        before=(window.preferences.source_language,window.preferences.target_language,window.preferences.device,
                window.text_page.source.editor.toPlainText(),settings.output_location(),window.preferences.domain)
        dialog._change_locale(LOCALES[locale])
        app.processEvents()
        assert (window.preferences.source_language,window.preferences.target_language,window.preferences.device,
                window.text_page.source.editor.toPlainText(),settings.output_location(),window.preferences.domain)==before
        assert dialog.windowTitle()==localization.text('Настройки TreeTranslate')
        assert window.file_page.languages.source_combo.currentText()=='Китайский'
        assert window.file_page.languages.source_combo.itemText(3)==localization.text(window.file_page.languages.source_combo._items[3])
        fresh=SettingsService(QSettings(path,QSettings.Format.IniFormat))
        assert saved_locale(fresh)==locale
        assert not hasattr(window.file_page,'domain_combo') and not hasattr(window.text_page,'mode')
        assert [dialog.device_combo.itemText(i) for i in range(3)]==['Auto','CPU','GPU']
    finally:dialog.close();window.close()


def test_catalog_coverage_and_templates():
    from tools.build_ui_catalogs import source_keys
    import re
    expected=set(source_keys())
    for locale in LOCALES:
        catalog=json.loads((ROOT/(locale+'.json')).read_text('utf-8'))
        assert expected<=catalog.keys(), (locale,expected-catalog.keys())
        for key,value in catalog.items():
            assert value and '<unk>' not in value
            assert sorted(re.findall(r'\{\d+\}',key))==sorted(re.findall(r'\{\d+\}',value))


def test_obsolete_settings_normalization_and_interface_persistence(tmp_path):
    app=QApplication.instance() or QApplication([])
    path=str(tmp_path/'settings.ini');settings=SettingsService(QSettings(path,QSettings.Format.IniFormat))
    for k,v in {'performance/mode':'Турбо','performance/device':'broken','performance/cpu_threads':'bad',
                'translation/domain':'metallurgy','general/ui_language':'unknown','interface/eta':False,
                'interface/detailed_progress':False}.items():settings.save_value(k,v)
    policy=settings.load_performance()
    assert policy.mode=='Автоматический' and policy.device=='Auto' and policy.cpu_threads=='Автоматически'
    assert saved_locale(settings)=='ru-RU'
    settings.sync()
    from app.gui.widgets.progress_panel import ProgressPanel
    panel=ProgressPanel();fresh=SettingsService(QSettings(path,QSettings.Format.IniFormat))
    panel.apply_settings(fresh)
    assert panel.current.isHidden() and panel.remaining.isHidden()
    assert not panel.bar.isHidden() and not panel.batch.isHidden()
    fresh.save_value('interface/eta',True);panel.apply_settings(fresh)
    assert not panel.remaining.isHidden() and panel.current.isHidden()
    fresh.save_value('interface/detailed_progress',True);panel.apply_settings(fresh)
    assert not panel.current.isHidden()
    panel.close()


def test_hardware_detection_without_cuda_keeps_device_choice(monkeypatch):
    from app.services.hardware_profile_service import HardwareProfile, HardwareProfileService
    from app.engine.runtime.device_manager import DeviceManager
    from test_document_service import wait_for
    app=QApplication.instance() or QApplication([])
    monkeypatch.setattr(HardwareProfileService,'detect',lambda self:HardwareProfile('Test CPU',8,16,'Other GPU',0))
    monkeypatch.setattr(HardwareProfileService,'save',lambda *a:None)
    monkeypatch.setattr(DeviceManager,'gpu_available',lambda self:False)
    dialog=SettingsDialog()
    try:
        dialog.device_combo.setCurrentText('GPU')
        dialog._show_hardware_recommendation()
        wait_for(lambda:dialog.detect_button.isEnabled())
        assert 'Test CPU' in dialog.hardware_value.text()
        assert 'CUDA: Недоступно' in dialog.hardware_value.text()
        assert 'VRAM: Не определено' in dialog.hardware_value.text()
        assert dialog.preferences.device=='GPU'
    finally:dialog.close()
