from hashlib import sha256
from types import SimpleNamespace
from threading import Thread, Event
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from app.documents.control import JobControl
from app.documents.docx_document import DocxDocument, validate_docx
from app.documents.errors import InvalidDocumentError, SourceChangedError
from app.documents.job import DocumentJob, DocumentConfig, safe_name, safe_filename_base
from app.documents.scanner import scan_sources
from app.engine.errors import TranslationCancelledError
from app.engine.types import DevicePreference, PerformanceProfile
from app.documents.errors import DocumentError


def fixture(path):
    doc = Document()
    doc.add_heading('Installation guide', 1)
    p = doc.add_paragraph()
    p.add_run('Hello ').bold = True
    p.add_run('world!').italic = True
    doc.add_paragraph('First step', style='List Bullet')
    doc.add_table(rows=1, cols=2).cell(0, 0).text = 'Table text'
    doc.sections[0].header.paragraphs[0].text = 'Header'
    doc.sections[0].footer.paragraphs[0].text = 'Footer'
    doc.add_picture(str(Path(__file__).resolve().parents[1] / 'assets/icons/languages/language_ru.png'))
    p = doc.add_paragraph('See https://example.org and ')
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to('https://example.org', RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = 'website'
    run.append(text)
    link.append(run)
    p._p.append(link)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'DATE')
    field_run = OxmlElement('w:r')
    field_text = OxmlElement('w:t')
    field_text.text = 'Protected date'
    field_run.append(field_text)
    field.append(field_run)
    p._p.append(field)
    doc.save(path)
    return path


def job(paths, config=None, control=None, translate=None, **callbacks):
    control = control or JobControl()
    return DocumentJob(scan_sources(paths, control).files, config or DocumentConfig(source='en'), control,
                       translate or (lambda request, cancel: SimpleNamespace(translated_text='Перевод ' + request.text)),
                       lambda text, source, target: (source, target), **callbacks)


def test_docx_preserves_package_and_translates_logical_runs(tmp_path):
    source = fixture(tmp_path / '原文.docx')
    before = source.read_bytes()
    seen = []
    progress = []
    def translate(request, cancel):
        seen.append(request.text)
        return SimpleNamespace(translated_text='Перевод ' + request.text)
    output, = job([source], translate=translate, progress=progress.append).run()
    assert source.read_bytes() == before
    assert output.name == '原文_ru.docx'
    assert 'Hello world!' in seen
    assert all('Protected date' not in text and 'https://' not in text for text in seen)
    assert {'Header', 'Footer', 'Table text'} <= set(seen)
    validate_docx(output)
    with ZipFile(source) as original, ZipFile(output) as result:
        assert original.namelist() == result.namelist()
        for name in original.namelist():
            if name not in {'word/document.xml', 'word/header1.xml', 'word/footer1.xml'}:
                assert original.read(name) == result.read(name)
        xml = result.read('word/document.xml')
        assert b'https://example.org' in xml and b'Protected date' in xml
        assert b'<w:b' in xml and b'<w:i' in xml and b'w:hyperlink' in xml
    assert progress[-1].percent == 100
    assert all(p.percent < 100 for p in progress[:-1])


def test_collision_and_no_original_overwrite(tmp_path):
    source = fixture(tmp_path / 'manual.docx')
    first, = job([source]).run()
    content = first.read_bytes()
    second, = job([source]).run()
    assert second.name == 'manual_ru (1).docx'
    assert first.read_bytes() == content
    third, = job([source], DocumentConfig(source='en', template='{name}')).run()
    assert third != source and third.name == 'manual (1).docx'


@pytest.mark.parametrize('text', ['', '12345', '你好，世界。'])
def test_empty_numeric_chinese(tmp_path, text):
    path = tmp_path / 'test.docx'
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)
    output, = job([path]).run()
    validate_docx(output)


def test_scan_recursive_deduplicates_and_skips(tmp_path):
    folder = tmp_path / 'nested'
    folder.mkdir()
    source = fixture(folder / 'test.DOCX')
    (folder / 'test.pdf').write_bytes(b'%PDF')
    result = scan_sources([tmp_path, source], JobControl())
    assert len(result.files) == 1 and len(result.skipped) == 1
    assert result.files[0].relative.parts == ('nested', 'test.DOCX')
    output, = job([folder], DocumentConfig(source='en', translate_directories=True)).run()
    assert output.parent.name == 'Перевод nested_ru'


def test_corrupt(tmp_path):
    source = tmp_path / 'bad.docx'
    source.write_bytes(b'bad')
    with pytest.raises(InvalidDocumentError):
        job([source]).run()
    assert list(tmp_path.iterdir()) == [source]


def test_cancel_no_final_or_temporary(tmp_path):
    source = fixture(tmp_path / 'source.docx')
    control = JobControl()
    before = sha256(source.read_bytes()).hexdigest()
    def translate(request, cancel):
        control.cancel()
        return SimpleNamespace(translated_text='Отмена')
    with pytest.raises(TranslationCancelledError):
        job([source], control=control, translate=translate).run()
    assert list(tmp_path.iterdir()) == [source]
    assert sha256(source.read_bytes()).hexdigest() == before


def test_changed_source_prevents_publication(tmp_path):
    source = fixture(tmp_path / 'source.docx')
    def translate(request, cancel):
        source.write_bytes(b'external modification')
        return SimpleNamespace(translated_text='Перевод')
    with pytest.raises(SourceChangedError):
        job([source], translate=translate).run()
    assert list(tmp_path.iterdir()) == [source]


def test_pause_resume_and_cancel():
    control = JobControl()
    control.pause()
    entered, exited = Event(), Event()
    def wait():
        entered.set()
        control.checkpoint()
        exited.set()
    thread = Thread(target=wait)
    thread.start()
    assert entered.wait(1) and not exited.wait(.05)
    control.resume()
    thread.join(1)
    assert exited.is_set()
    control.pause()
    control.cancel()
    with pytest.raises(TranslationCancelledError):
        control.checkpoint()


@pytest.mark.parametrize('name', ['CON', 'a/b:*?', '..', 'NUL.docx'])
def test_safe_names(name):
    assert safe_name(name) not in {'CON', '.', '..', 'NUL.docx'}
    assert not any(c in safe_name(name) for c in '/\\:*?')


def test_cancel_during_validation_removes_temp(tmp_path, monkeypatch):
    source = fixture(tmp_path / 'source.docx')
    control = JobControl()
    import app.documents.job as module
    validate = module.validate_docx
    def cancelling_validate(path, structure):
        validate(path, structure)
        control.cancel()
    monkeypatch.setattr(module, 'validate_docx', cancelling_validate)
    with pytest.raises(TranslationCancelledError):
        job([source], control=control).run()
    assert list(tmp_path.iterdir()) == [source]


def test_failed_validation_removes_temp(tmp_path, monkeypatch):
    source = fixture(tmp_path / 'source.docx')
    before = source.read_bytes()
    def fail(*args):
        raise InvalidDocumentError()
    monkeypatch.setattr('app.documents.job.validate_docx', fail)
    with pytest.raises(InvalidDocumentError):
        job([source]).run()
    assert list(tmp_path.iterdir()) == [source] and source.read_bytes() == before


def test_directory_structure_and_translation_collisions(tmp_path):
    root = tmp_path / 'source'
    for name in ['one', 'two']:
        child = root / name
        child.mkdir(parents=True)
        fixture(child / 'manual.docx')
    config = DocumentConfig(source='en', output=tmp_path / 'output', translate_directories=True)
    results = job([root], config, translate=lambda r, c: SimpleNamespace(translated_text='Имя')).run()
    assert len(results) == 2
    assert {p.parent.name for p in results} == {'Имя', 'Имя (1)'}
    assert len({p.parent.parent for p in results}) == 1


def test_unsupported_pdf_and_auto_target(tmp_path):
    pdf = tmp_path / 'test.pdf'
    pdf.write_bytes(b'%PDF')
    with pytest.raises(DocumentError, match='Формат пока не поддерживается'):
        job([pdf]).run()
    source = fixture(tmp_path / 'source.docx')
    with pytest.raises(DocumentError, match='явно'):
        job([source], DocumentConfig(target='auto')).run()


def test_no_translatable_text_never_loads_detector_or_model(tmp_path):
    source = tmp_path / 'empty.docx'
    Document().save(source)
    work = job([source])
    def forbidden(*args):
        raise AssertionError('No text must not load models')
    work.resolve = work.translate = forbidden
    output, = work.run()
    assert source.read_bytes() == output.read_bytes()


def test_directory_selection_preserves_all_original_hashes(tmp_path):
    root = tmp_path / 'source'
    root.mkdir()
    paths = [fixture(root / (name + '.docx')) for name in ['one', '第二']]
    before = {p: sha256(p.read_bytes()).hexdigest() for p in paths}
    outputs = job([root]).run()
    assert len(outputs) == 2
    assert all(sha256(p.read_bytes()).hexdigest() == digest for p, digest in before.items())


def test_sample_includes_end_of_large_document(tmp_path):
    source = tmp_path / 'long.docx'
    document = Document()
    for i in range(100):
        document.add_paragraph(f'Paragraph {i} ' + 'text ' * 200)
    document.save(source)
    sample = DocxDocument(source).sample()
    assert 'Paragraph 0 ' in sample and 'Paragraph 99 ' in sample
    assert len(sample) <= 8000


def test_opaque_binary_and_nested_tables_preserved(tmp_path):
    source = fixture(tmp_path / 'objects.docx')
    document = Document(source)
    document.tables[0].cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = 'Nested table'
    document.save(source)
    with ZipFile(source, 'a') as archive:
        archive.writestr('word/embeddings/opaque.bin', b'\x00\x01opaque binary\xff')
    output, = job([source]).run()
    assert 'Перевод Nested table' in Document(output).tables[0].cell(0, 1).tables[0].cell(0, 0).text
    with ZipFile(source) as original, ZipFile(output) as translated:
        for name in original.namelist():
            if name.startswith(('word/media/', 'word/embeddings/')):
                assert original.read(name) == translated.read(name)


def test_private_text_not_logged(tmp_path, caplog):
    import logging
    source = tmp_path / 'private.docx'
    doc = Document()
    doc.add_paragraph('secret document content 932761')
    doc.save(source)
    with caplog.at_level(logging.INFO):
        job([source]).run()
    assert 'secret document content' not in caplog.text
    assert 'Перевод' not in caplog.text
    assert 'segments=1' in caplog.text


def filename_translator(mapping, calls=None):
    def translate(request, cancel):
        if calls is not None:
            calls.append(request)
        value = mapping.get(request.text, 'Перевод ' + request.text)
        return SimpleNamespace(translated_text=value)
    return translate


@pytest.mark.parametrize(
    ('source_name', 'translated', 'expected'),
    [
        ('hi.docx', 'Привет', 'Привет_ru.docx'),
        ('使用说明.docx', 'Руководство', 'Руководство_ru.docx'),
    ],
)
def test_filename_translation_uses_stem_only_and_preserves_source(tmp_path, source_name, translated, expected):
    source = fixture(tmp_path / source_name)
    before = sha256(source.read_bytes()).hexdigest()
    config = DocumentConfig(source='en', translate_filenames=True)
    output, = job([source], config, translate=filename_translator({source.stem: translated})).run()
    assert output.name == expected
    assert output.suffix == source.suffix
    assert sha256(source.read_bytes()).hexdigest() == before


def test_filename_translation_toggle_off_keeps_original_base(tmp_path):
    source = fixture(tmp_path / 'hi.docx')
    calls = []
    output, = job(
        [source], DocumentConfig(source='en', translate_filenames=False),
        translate=filename_translator({'hi': 'Привет'}, calls),
    ).run()
    assert output.name == 'hi_ru.docx'
    assert all(request.text != 'hi' for request in calls)


@pytest.mark.parametrize(
    ('translated', 'expected'),
    [
        ('bad<>:"/\\|?*name. ', 'bad_________name_ru.docx'),
        ('CON', 'original_ru.docx'),
        ('   ...   ', 'original_ru.docx'),
        ('', 'original_ru.docx'),
    ],
)
def test_translated_filename_is_windows_safe_or_falls_back(tmp_path, translated, expected):
    source = fixture(tmp_path / 'original.docx')
    output, = job(
        [source], DocumentConfig(source='en', translate_filenames=True),
        translate=filename_translator({'original': translated}),
    ).run()
    assert output.name == expected
    assert not any(character in output.stem for character in '<>:"/\\|?*')


def test_translated_filename_collision_uses_duplicate_safe_name(tmp_path):
    source = fixture(tmp_path / 'hi.docx')
    config = DocumentConfig(source='en', translate_filenames=True)
    translate = filename_translator({'hi': 'Привет'})
    first, = job([source], config, translate=translate).run()
    second, = job([source], config, translate=translate).run()
    assert first.name == 'Привет_ru.docx'
    assert second.name == 'Привет_ru (1).docx'


def test_nested_batch_translates_each_filename_independently(tmp_path):
    root = tmp_path / 'documents'
    nested = root / 'nested'
    nested.mkdir(parents=True)
    first = fixture(root / 'hi.docx')
    second = fixture(nested / '使用说明.docx')
    before = {path: sha256(path.read_bytes()).hexdigest() for path in (first, second)}
    outputs = job(
        [root],
        DocumentConfig(source='en', output=tmp_path / 'results', translate_filenames=True),
        translate=filename_translator({'hi': 'Привет', '使用说明': 'Руководство'}),
    ).run()
    assert {path.name for path in outputs} == {'Привет_ru.docx', 'Руководство_ru.docx'}
    assert len({path.parent for path in outputs}) == 2
    assert all(sha256(path.read_bytes()).hexdigest() == digest for path, digest in before.items())


def test_directory_and_filename_translation_toggles_are_independent(tmp_path):
    root = tmp_path / 'Folder'
    root.mkdir()
    source = fixture(root / 'hi.docx')
    translate = filename_translator({'Folder': 'Папка', 'hi': 'Привет'})

    filename_only, = job(
        [root],
        DocumentConfig(source='en', output=tmp_path / 'filename-output',
                       translate_directories=False, translate_filenames=True),
        translate=translate,
    ).run()
    directory_only, = job(
        [root],
        DocumentConfig(source='en', output=tmp_path / 'directory-output',
                       translate_directories=True, translate_filenames=False),
        translate=translate,
    ).run()
    assert filename_only.name == 'Привет_ru.docx'
    assert filename_only.parent.name == 'Folder_ru'
    assert directory_only.name == 'hi_ru.docx'
    assert directory_only.parent.name == 'Папка_ru'


def test_scanner_skips_office_temp_hidden_and_service_docx(tmp_path):
    valid = fixture(tmp_path / 'valid.docx')
    ignored = [fixture(tmp_path / name) for name in ('~$draft.docx', '.hidden.docx', '.treetranslate-work.docx')]
    result = scan_sources([tmp_path], JobControl())
    assert tuple(item.path for item in result.files) == (valid.resolve(),)
    assert {path.resolve() for path in result.skipped} == {path.resolve() for path in ignored}


def test_safe_filename_base_rejects_windows_device_names_and_trailing_dots():
    assert safe_filename_base('LPT9. ', 'manual') == 'manual'
    assert safe_filename_base('valid. ', 'manual') == 'valid'


def test_filename_request_reuses_selected_language_device_profile_and_threads(tmp_path):
    source = fixture(tmp_path / 'hi.docx')
    calls = []
    config = DocumentConfig(
        source='zh', target='ru', device=DevicePreference.GPU,
        profile=PerformanceProfile.TURBO, threads=4, translate_filenames=True,
    )
    output, = job(
        [source], config,
        translate=filename_translator({'hi': 'Привет'}, calls),
    ).run()
    filename_request = next(request for request in calls if request.text == 'hi')
    assert filename_request.source_language == 'zh'
    assert filename_request.target_language == 'ru'
    assert filename_request.device_preference is DevicePreference.GPU
    assert filename_request.performance_profile is PerformanceProfile.TURBO
    assert filename_request.cpu_threads == 4
    assert output.name == 'Привет_ru.docx'
