from dataclasses import replace
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import struct
import stat
from zipfile import ZipFile, ZipInfo, ZIP_STORED

import pytest
from docx import Document

from app.documents.control import JobControl
from app.documents.errors import DocumentError
from app.documents.job import DocumentJob, DocumentConfig
from app.documents.scanner import scan_sources
from app.documents.zip_archive import ArchiveLimits, inventory, validate_output
from app.engine.errors import TranslationCancelledError


def docx_bytes(text='Source text'):
    doc = Document(); doc.add_paragraph(text)
    stream = BytesIO(); doc.save(stream)
    return stream.getvalue()


def archive(tmp_path, entries, name='车身尺寸.zip'):
    path = tmp_path / name
    with ZipFile(path, 'w') as z:
        for entry, data in entries:
            z.writestr(entry, data)
    return path


def run(path, config=None, control=None, translator=None, completed=None, progress=None):
    control = control or JobControl()
    files = scan_sources([path], control).files
    config = config or DocumentConfig(source='zh', target='ru', translate_filenames=True, translate_directories=True)
    names = {'车身尺寸': 'Размеры кузова', '内部': 'Внутренняя часть', '前车身': 'Передняя часть кузова',
             '维修': 'Ремонт', '车身维修': 'Ремонт кузова'}
    translator = translator or (lambda request, cancel: SimpleNamespace(translated_text=names.get(request.text, 'Перевод текста')))
    return DocumentJob(files, config, control, translator, lambda text, source, target: ('zh', target),
        progress=progress.append if progress is not None else lambda p: None,
        file_completed=lambda s, o: completed.append((s, o)) if completed is not None else None).run()


def test_zip_to_zip_unicode_structure_copy_and_source_immutability(tmp_path):
    nested = BytesIO()
    with ZipFile(nested, 'w') as z:
        z.writestr('../unsafe.txt', b'nested archive must not be opened')
    source = archive(tmp_path, [('内部/', b''), ('内部/车身维修.docx', docx_bytes()),
        ('前车身/车身维修.docx', docx_bytes()), ('空目录/', b''),
        ('内部/image.png', b'image-bytes'), ('内部/nested.zip', nested.getvalue()), ('ä_日本語_é.txt', b'raw bytes')])
    before = sha256(source.read_bytes()).hexdigest()
    progress, completed = [], []
    output, = run(source, completed=completed, progress=progress)
    assert output.name == 'Размеры кузова_ru.zip'
    assert output.is_file() and not any(p.is_dir() for p in tmp_path.iterdir())
    assert sha256(source.read_bytes()).hexdigest() == before
    assert all(p.percent < 100 for p in progress[:-1]) and progress[-1].percent == 100
    assert progress[-1].processed == progress[-1].total and progress[-1].total > 0
    assert len(completed) == 2 and all(o == output for s, o in completed)
    with ZipFile(output) as z:
        names = z.namelist()
        assert 'Внутренняя часть/Ремонт кузова_ru.docx' in names
        assert 'Передняя часть кузова/Ремонт кузова_ru.docx' in names
        assert z.read('Внутренняя часть/image.png') == b'image-bytes'
        assert z.read('Внутренняя часть/nested.zip') == nested.getvalue()
        assert z.read('ä_日本語_é.txt') == b'raw bytes'
        assert z.testzip() is None
        assert all(i.flag_bits & 0x800 for i in z.infolist() if not i.filename.isascii())
        assert any(n.endswith('/') and '空目录' not in n for n in names)
        assert not any(n.startswith('Размеры кузова/') for n in names)
        for info in z.infolist():
            if info.filename.endswith('.docx'):
                assert Document(BytesIO(z.read(info))).paragraphs[0].text == 'Перевод текста'


def test_translated_name_extensions_and_internal_external_collisions(tmp_path):
    source = archive(tmp_path, [('甲/one.pdf.pdf.docx', docx_bytes()), ('乙/two.docx', docx_bytes()),
                              ('root-a.docx', docx_bytes()), ('root-b.docx', docx_bytes())])
    calls = []
    def translate(request, cancel):
        calls.append(request.text)
        return SimpleNamespace(translated_text='Одинаково.docx' if '.' in request.text or request.text.startswith('root') else 'Одинаково')
    first, = run(source, translator=translate)
    second, = run(source, translator=translate)
    assert first != second and second.stem.endswith('(1)')
    with ZipFile(first) as z:
        assert len(set(n.casefold() for n in z.namelist())) == len(z.namelist())
        docs = [n for n in z.namelist() if n.endswith('.docx')]
        assert len(docs) == 4 and all(not n.endswith('.docx.docx') for n in docs)
        assert any('(1)' in n for n in docs)
    assert 'root-a.docx' not in calls


def test_folder_component_translated_once_and_name_flags_off(tmp_path):
    source = archive(tmp_path, [('one/内部/a.docx', docx_bytes()), ('two/内部/b.docx', docx_bytes())])
    calls = []
    def translate(request, cancel):
        calls.append(request.text)
        return SimpleNamespace(translated_text='Перевод')
    run(source, translator=translate)
    assert calls.count('内部') == 1
    output, = run(source, config=DocumentConfig(source='zh', translate_directories=False, translate_filenames=False))
    with ZipFile(output) as z:
        assert 'one/内部/a_ru.docx' in z.namelist()
        assert 'two/内部/b_ru.docx' in z.namelist()
    assert output.name == '车身尺寸_ru.zip'


@pytest.mark.parametrize('name', ['../escape.docx', '/absolute.docx', 'C:/bad.docx', 'a\\b.docx',
    'NUL.docx', 'folder/CON.txt', 'folder./x.docx', 'a//x.docx', 'a/./x.docx', 'a/../x.docx', 'a:stream.docx'])
def test_unsafe_paths_rejected(tmp_path, name):
    source = archive(tmp_path, [(name, b'bad')])
    if '\\' in name:
        # ZipInfo normalizes platform separators while writing on Windows.
        source.write_bytes(source.read_bytes().replace(name.replace('\\', '/').encode(), name.encode()))
    with pytest.raises(DocumentError):
        scan_sources([source], JobControl())
    assert len(list(tmp_path.iterdir())) == 1


@pytest.mark.parametrize('entries', [[('same.docx', b'a'), ('SAME.docx', b'b')], [('a', b'a'), ('a/b', b'b')]])
def test_duplicates_and_file_directory_conflicts_rejected(tmp_path, entries):
    source = archive(tmp_path, entries)
    with pytest.raises(DocumentError):
        scan_sources([source], JobControl())


def test_symlink_encrypted_corrupt_and_resource_limits(tmp_path):
    source = tmp_path / 'link.zip'
    link = ZipInfo('link'); link.create_system = 3; link.external_attr = (stat.S_IFLNK | 0o777) << 16
    with ZipFile(source, 'w') as z:
        z.writestr(link, b'../target')
    with pytest.raises(DocumentError): inventory(source, JobControl())
    source = archive(tmp_path, [('a.txt', b'abc')])
    data = bytearray(source.read_bytes())
    local = data.index(b'PK\x03\x04'); central = data.index(b'PK\x01\x02')
    struct.pack_into('<H', data, local+6, 1); struct.pack_into('<H', data, central+8, 1)
    source.write_bytes(data)
    with pytest.raises(DocumentError, match='Зашифрованные'): inventory(source, JobControl())
    source.write_bytes(b'broken')
    with pytest.raises(DocumentError): inventory(source, JobControl())
    source = archive(tmp_path, [('a.txt', b'abc'), ('b.txt', b'xyz')])
    for limits in (ArchiveLimits(members=1), ArchiveLimits(total_bytes=4), ArchiveLimits(member_bytes=2)):
        with pytest.raises(DocumentError): inventory(source, JobControl(), limits)
    bomb = tmp_path / 'bomb.zip'
    from zipfile import ZIP_DEFLATED
    with ZipFile(bomb, 'w', compression=ZIP_DEFLATED) as z: z.writestr('bomb.txt', b'0'*100000)
    with pytest.raises(DocumentError): inventory(bomb, JobControl())


@pytest.mark.parametrize('cancel', [False, True])
def test_document_failure_preserves_sources_cancel_publishes_nothing_cleanup(tmp_path, monkeypatch, cancel):
    import app.documents.archive_job as module
    workspaces = []
    original = module.TemporaryDirectory
    def temporary(*args, **kwargs):
        workspace = original(*args, **kwargs); workspaces.append(Path(workspace.name)); return workspace
    monkeypatch.setattr(module, 'TemporaryDirectory', temporary)
    source = archive(tmp_path, [('内部/a.docx', docx_bytes()), ('内部/b.docx', docx_bytes())])
    before = source.read_bytes()
    control = JobControl()
    def translate(request, event):
        if request.text == 'Source text':
            if cancel:
                control.cancel(); control.checkpoint()
            raise DocumentError('Expected translation error')
        return SimpleNamespace(translated_text='Имя')
    if cancel:
        with pytest.raises(TranslationCancelledError):
            run(source, control=control, translator=translate)
    else:
        output, = run(source, control=control, translator=translate)
        with ZipFile(output) as final,ZipFile(source) as original_zip:
            for member in original_zip.namelist():
                assert final.read(member)==original_zip.read(member)
    assert source.read_bytes() == before
    assert all(not p.exists() for p in workspaces)
    if cancel:
        assert list(tmp_path.iterdir()) == [source]


def test_crc_output_validation_and_no_success_before_validation(tmp_path, monkeypatch):
    source = archive(tmp_path, [('a.docx', docx_bytes())])
    with pytest.raises(DocumentError): validate_output(source, ['missing.docx'], JobControl())
    data = bytearray(source.read_bytes())
    # Stored member: alter payload while keeping its CRC intact.
    offset = 30 + struct.unpack_from('<H', data, 26)[0] + struct.unpack_from('<H', data, 28)[0]
    data[offset+10] ^= 1
    source.write_bytes(data)
    with pytest.raises(DocumentError): inventory(source, JobControl())
    source = archive(tmp_path, [('a.docx', docx_bytes())])
    import app.documents.archive_job as module
    monkeypatch.setattr(module, 'validate_output', lambda *args: (_ for _ in ()).throw(DocumentError('Expected output error')))
    completed = []
    with pytest.raises(DocumentError): run(source, completed=completed)
    assert not completed and list(tmp_path.iterdir()) == [source]


def test_zip_with_only_assets_still_returns_complete_zip(tmp_path):
    source = archive(tmp_path, [('readme.txt', b'unchanged')])
    output, = run(source)
    with ZipFile(output) as z: assert z.read('readme.txt') == b'unchanged'


def test_cancel_during_pack_and_publication_collision_race(tmp_path, monkeypatch):
    source = archive(tmp_path, [('a.docx', docx_bytes())])
    control = JobControl()
    class Progress(list):
        def append(self, progress):
            super().append(progress)
            if progress.stage == 'ARCHIVE_PACKING':
                control.cancel()
    with pytest.raises(TranslationCancelledError):
        run(source, control=control, progress=Progress())
    assert list(tmp_path.iterdir()) == [source]
    from app.documents.zip_archive import publish
    temporary = tmp_path / 'temporary.zip'
    temporary.write_bytes(b'validated')
    destination = tmp_path / 'result.zip'
    control = JobControl()
    original = control.publish
    calls = []
    def race(action):
        if not calls:
            destination.write_bytes(b'concurrent user output')
        calls.append(action)
        return original(action)
    control.publish = race
    result = publish(temporary, destination, source, sha256(source.read_bytes()).hexdigest(), control)
    assert result.name == 'result (1).zip'
    assert destination.read_bytes() == b'concurrent user output'
    assert result.read_bytes() == b'validated'


def test_source_changed_after_scan_prevents_output_and_no_auto_target(tmp_path):
    source = archive(tmp_path, [('a.docx', docx_bytes())])
    control = JobControl()
    files = scan_sources([source], control).files
    archive(tmp_path, [('a.docx', docx_bytes('Changed text'))])
    with pytest.raises(DocumentError, match='Исходный файл изменился'):
        DocumentJob(files, DocumentConfig(source='zh'), control, lambda *a: None, lambda *a: ('zh', 'ru')).run()
    with pytest.raises(DocumentError, match='язык перевода явно'):
        run(source, config=DocumentConfig(target='auto'))
    assert list(tmp_path.iterdir()) == [source]


def test_numeric_archive_name_and_compressible_assets(tmp_path):
    source = archive(tmp_path, [('zeros.bin', b'0'*100000)], name='123.zip')
    output, = run(source, config=DocumentConfig(translate_filenames=True))
    assert output.name == '123_ru.zip'
    with ZipFile(output) as z:
        assert z.read('zeros.bin') == b'0'*100000


def test_translated_folder_does_not_replace_copied_asset_and_unicode_alias_collisions(tmp_path):
    source = archive(tmp_path, [('目录/a.docx', docx_bytes()), ('image.png', b'asset'),
                               ('first.docx', docx_bytes()), ('second.docx', docx_bytes())])
    names = {'目录': 'image.png', 'first': 'Straße', 'second': 'STRASSE'}
    def translate(request, cancel):
        return SimpleNamespace(translated_text=names.get(request.text, 'Перевод'))
    output, = run(source, translator=translate)
    with ZipFile(output) as z:
        assert z.read('image.png') == b'asset'
        assert 'image.png (1)/Перевод_ru.docx' in z.namelist()
        assert 'Straße_ru.docx' in z.namelist()
        assert 'STRASSE_ru (1).docx' in z.namelist()


def test_pause_and_resume_archive_before_pack(tmp_path):
    from threading import Event, Thread
    source = archive(tmp_path, [('a.docx', docx_bytes())])
    control = JobControl()
    paused, outputs, errors = Event(), [], []
    class Progress(list):
        def append(self, progress):
            if progress.stage == 'ARCHIVE_PACKING':
                control.pause()
                paused.set()
    def worker():
        try:
            outputs.extend(run(source, control=control, progress=Progress()))
        except Exception as error:
            errors.append(error)
    thread = Thread(target=worker)
    thread.start()
    try:
        assert paused.wait(5)
        assert not outputs
        control.resume()
        thread.join(5)
        assert not thread.is_alive() and not errors and len(outputs) == 1
        with ZipFile(outputs[0]) as z: assert z.testzip() is None
    finally:
        control.cancel(); thread.join(5)


def test_mixed_folder_and_zip_keep_one_output_folder(tmp_path):
    folder = tmp_path / 'manuals'; folder.mkdir()
    for name in ('first.docx', 'second.docx'):
        (folder / name).write_bytes(docx_bytes())
    source = archive(tmp_path, [('inside.docx', docx_bytes())])
    control = JobControl()
    files = scan_sources([folder, source], control).files
    outputs = DocumentJob(files, DocumentConfig(source='zh', translate_directories=True), control,
        lambda request, event: SimpleNamespace(translated_text='Перевод'),
        lambda text, source, target: ('zh', 'ru')).run()
    documents = [p for p in outputs if p.suffix == '.docx']
    assert len(documents) == 2 and documents[0].parent == documents[1].parent
    assert len([p for p in outputs if p.suffix == '.zip']) == 1
