"""Optional real-model DOCX coverage; regular fixtures never need model weights."""
from hashlib import sha256
from pathlib import Path
import socket

from docx import Document
import pytest

from app.documents.control import JobControl
from app.documents.docx_document import validate_docx
from app.documents.job import DocumentConfig, DocumentJob
from app.documents.scanner import scan_sources
from app.engine.factory import create_translation_engine
from app.engine.types import DevicePreference, PerformanceProfile

pytestmark = pytest.mark.integration


@pytest.fixture(scope='module')
def engine():
    if not all((Path('vendor/models') / p).is_file() for p in (
        'argos/argos-en-ru/model/model.bin', 'm2m100-418m-int8/model/model.bin')):
        pytest.skip('Prepared local Argos/M2M100 models are absent')
    engine = create_translation_engine()
    yield engine
    engine.shutdown()


@pytest.mark.parametrize('device', list(DevicePreference))
@pytest.mark.parametrize('text,language', [('Save the configuration file before restarting the application.', 'en'),
                                         ('请在重新启动应用程序之前保存配置文件。', 'zh')])
def test_real_docx(engine, tmp_path, device, text, language, monkeypatch):
    if device == DevicePreference.GPU and not engine.devices.gpu_available():
        pytest.skip('No CUDA device')
    path = tmp_path / '原文.docx'
    doc = Document()
    doc.add_paragraph(text)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = text
    doc.sections[0].header.paragraphs[0].text = text
    doc.sections[0].footer.paragraphs[0].text = text
    doc.save(path)
    digest = sha256(path.read_bytes()).hexdigest()
    attempts = []
    def forbidden(*args, **kwargs):
        attempts.append(True)
        raise AssertionError('Network attempted')
    monkeypatch.setattr(socket, 'getaddrinfo', forbidden)
    monkeypatch.setattr(socket, 'create_connection', forbidden)
    control = JobControl()
    requests, resolved = [], []
    def translate(request, cancelled):
        requests.append(request)
        return engine.translate(request, cancelled)
    def resolve(text, source, target):
        pair = engine.languages.resolve(text, source, target)
        resolved.append(pair)
        return pair
    progress = []
    with engine.runtime.keep_warm():
        output, = DocumentJob(scan_sources([path], control).files,
                              DocumentConfig(device=device, profile=PerformanceProfile.BALANCED), control,
                              translate, resolve, progress.append).run()
    assert resolved == [(language, 'ru')]
    assert all(r.source_language == language for r in requests)
    assert sha256(path.read_bytes()).hexdigest() == digest
    assert not attempts
    validate_docx(output)
    result = Document(output)
    assert result.paragraphs[0].text != text
    assert all(any('а' <= c.lower() <= 'я' for c in value) for value in (
        result.paragraphs[0].text, result.tables[0].cell(0, 0).text,
        result.sections[0].header.paragraphs[0].text, result.sections[0].footer.paragraphs[0].text))
    assert progress[-1].percent == 100


@pytest.mark.parametrize(
    ('filename', 'text', 'language'),
    [
        ('hi.docx', 'Hello world.', 'en'),
        ('使用说明.docx', '请保存此文件。', 'zh'),
    ],
)
def test_real_filename_translation_uses_existing_router(engine, tmp_path, filename, text, language):
    source = tmp_path / filename
    document = Document()
    document.add_paragraph(text)
    document.save(source)
    digest = sha256(source.read_bytes()).hexdigest()
    control = JobControl()
    with engine.runtime.keep_warm():
        output, = DocumentJob(
            scan_sources([source], control).files,
            DocumentConfig(source=language, device=DevicePreference.CPU,
                           profile=PerformanceProfile.BALANCED, translate_filenames=True),
            control, engine.translate, engine.languages.resolve,
        ).run()
    assert output.suffix == source.suffix
    assert output.stem.endswith('_ru')
    assert output.stem != f'{source.stem}_ru'
    assert any('а' <= character.lower() <= 'я' for character in output.stem)
    assert sha256(source.read_bytes()).hexdigest() == digest
