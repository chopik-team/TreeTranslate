from dataclasses import replace
from hashlib import sha256
from pathlib import Path
from types import SimpleNamespace
from threading import Event, Thread

import pytest
import pypdfium2 as pdfium
from pypdf import PdfReader, PdfWriter

from app.documents.control import JobControl
from app.documents.job import DocumentJob, DocumentConfig
from app.documents.scanner import scan_sources
from app.documents.pdf_document import PdfDocument
from app.documents.pdf_fonts import FontResolver
from app.documents.pdf_types import PdfKind, PdfError, PdfLimits
from app.engine.errors import TranslationCancelledError
from tools.pdf_fixtures import make_pdf, render


def pdf_job(paths, config=None, control=None, translate=None, **callbacks):
    control = control or JobControl()
    return DocumentJob(scan_sources(paths, control).files, config or DocumentConfig(source='en'), control,
                       translate or (lambda request, cancel: SimpleNamespace(translated_text='Перевод документа')),
                       lambda text, source, target: ('en', target), **callbacks)


def extract(path):
    with pdfium.PdfDocument(path) as document:
        texts = []
        for index in range(len(document)):
            page = document[index]
            textpage = page.get_textpage()
            texts.append(textpage.get_text_range())
            textpage.close()
            page.close()
        return '\n'.join(texts)


@pytest.mark.parametrize('text', ['Save the configuration file.', 'Сохраните настройки приложения.', '请保存配置文件，然后重新启动应用程序。'])
def test_real_pdf_objects_replaced_original_immutable(tmp_path, text):
    source = make_pdf(tmp_path / '原文.pdf', text, pages=2, link=True)
    before = source.read_bytes()
    progress = []
    output, = pdf_job([source], progress=progress.append).run()
    assert source.read_bytes() == before
    assert sha256(source.read_bytes()).digest() == sha256(before).digest()
    assert output.suffix == '.pdf'
    translated = extract(output)
    assert 'Перевод документа' in translated and text not in translated
    original, result = PdfReader(source), PdfReader(output)
    assert len(result.pages) == len(original.pages) == 2
    for old, new in zip(original.pages, result.pages):
        assert old.mediabox == new.mediabox
        assert new['/Annots'][0].get_object()['/A']['/URI'] == 'https://example.org/manual?id=42'
    assert progress[-1].percent == 100
    assert {'EXTRACTING', 'TRANSLATING', 'WRITING', 'VALIDATING', 'COMPLETED'} <= {p.stage for p in progress}


def test_fragments_become_logical_text(tmp_path):
    text = 'Save the configuration file.'
    source = make_pdf(tmp_path / 'fragments.pdf', text, fragments=True)
    doc = PdfDocument(source)
    assert any(s.text == text for s in doc.segments)
    assert len(doc.segments) == 2


def test_two_column_reading_order(tmp_path):
    source = make_pdf(tmp_path / 'columns.pdf', 'Text paragraph.', columns=True)
    segments = PdfDocument(source).segments
    left = [s.reading_order for s in segments if s.bbox[0] < 300 and s.text != 'User guide']
    right = [s.reading_order for s in segments if s.bbox[0] > 300]
    assert left and right and max(left) < min(right)


def test_mixed_preserves_image_and_vector_pixels(tmp_path):
    source = make_pdf(tmp_path / 'mixed.pdf', image=True)
    doc = PdfDocument(source)
    assert doc.classification == PdfKind.MIXED_PDF
    notices = []
    output, = pdf_job([source], warning=notices.append).run()
    assert any('изображений' in w for w in notices)
    old = render(source, tmp_path / 'before')[0]
    new = render(output, tmp_path / 'after')[0]
    from PIL import Image, ImageChops
    a, b = Image.open(old), Image.open(new)
    # This crop contains the image and colored vector background, no translated text.
    box = (600, 300, 850, 550)
    assert ImageChops.difference(a.crop(box), b.crop(box)).getbbox() is None


@pytest.mark.parametrize('options,kind', [({'empty': True, 'background': False}, PdfKind.EMPTY_PDF),
                                         ({'empty': True, 'image': True}, PdfKind.IMAGE_ONLY_PDF),
                                         ({'text': 'xy', 'image': True, 'lines': 1}, PdfKind.IMAGE_ONLY_PDF)])
def test_no_fake_translation_without_text(tmp_path, options, kind):
    source = make_pdf(tmp_path / 'no-text.pdf', **options)
    before = source.read_bytes()
    with pytest.raises(PdfError) as error:
        pdf_job([source]).run()
    assert error.value.kind == kind
    assert source.read_bytes() == before
    assert list(tmp_path.glob('*_ru.pdf')) == []


def test_corrupt_and_encrypted(tmp_path):
    bad = tmp_path / 'bad.pdf'
    bad.write_bytes(b'%PDF-1.7\ninvalid')
    with pytest.raises(PdfError) as error:
        PdfDocument(bad)
    assert error.value.kind == PdfKind.CORRUPTED_PDF
    source = make_pdf(tmp_path / 'source.pdf')
    encrypted = tmp_path / 'encrypted.pdf'
    writer = PdfWriter(clone_from=source)
    writer.encrypt('password', algorithm='RC4-128')
    writer.write(encrypted)
    with pytest.raises(PdfError) as error:
        PdfDocument(encrypted)
    assert error.value.kind == PdfKind.ENCRYPTED_PDF


@pytest.mark.parametrize('rotation', [90, 180, 270])
def test_page_rotation_preserved(tmp_path, rotation):
    source = make_pdf(tmp_path / 'rotated.pdf', rotation=rotation)
    output, = pdf_job([source]).run()
    with pdfium.PdfDocument(output) as document:
        page = document[0]
        assert page.get_rotation() == rotation
        page.close()
    assert 'Перевод документа' in extract(output)


def test_rotated_text_block(tmp_path):
    source = make_pdf(tmp_path / 'text-rotation.pdf', rotated_text=True)
    assert any(s.rotation == 90 for s in PdfDocument(source).segments)
    output, = pdf_job([source]).run()
    assert 'Перевод документа' in extract(output)


def test_full_overflow_text_in_continuation(tmp_path):
    source = make_pdf(tmp_path / 'small.pdf', 'Short text.', tiny=True, lines=1)
    full = 'Очень длинный перевод, который полностью сохранён. ' * 200
    notices = []
    output, = pdf_job([source], translate=lambda r, c: SimpleNamespace(translated_text=full), warning=notices.append).run()
    result = PdfReader(output)
    assert len(result.pages) > 1
    assert not any(p.get('/Annots') for p in result.pages)
    assert any('продолжении' in w for w in notices)
    import re
    body = re.sub(r'Продолжение\s+перевода|Страница\s*\d+\s*·\s*блок\s*\d+', '', extract(output))
    extracted = ''.join(body.split())
    assert extracted.count('Оченьдлинныйперевод,которыйполностьюсохранён.') >= 400


def test_url_email_ids_not_sent_to_engine(tmp_path):
    source = make_pdf(tmp_path / 'protected.pdf', 'Visit https://example.org or help@example.org ID-12345', lines=1)
    calls = []
    def translate(request, cancel):
        calls.append(request.text)
        return SimpleNamespace(translated_text='Перевод')
    output, = pdf_job([source], translate=translate).run()
    assert all('https://' not in t and '@' not in t and 'ID-12345' not in t for t in calls)
    result = PdfReader(output)
    contents = extract(output) + ''.join(a.get_object().get('/Contents', '') for a in result.pages[0].get('/Annots', []))
    assert 'https://example.org' in contents and 'help@example.org' in contents and 'ID-12345' in contents


def test_filename_directory_collision_and_mixed_job(tmp_path):
    from docx import Document
    root = tmp_path / 'Guides'
    source = make_pdf(root / 'Manual.pdf')
    doc = Document()
    doc.add_paragraph('Save the file.')
    doc.save(root / 'Other.docx')
    config = DocumentConfig(source='en', translate_filenames=True, translate_directories=True)
    outputs = pdf_job([root], config).run()
    assert {p.suffix for p in outputs} == {'.docx', '.pdf'}
    assert all(p.parent.name == 'Перевод документа_ru' for p in outputs)
    assert outputs[0].stem == 'Перевод документа_ru'
    first, = pdf_job([source]).run()
    second, = pdf_job([source]).run()
    assert first.name == 'Manual_ru.pdf' and second.name == 'Manual_ru (1).pdf'


def test_cancel_and_validation_failure_leave_no_temp(tmp_path, monkeypatch):
    source = make_pdf(tmp_path / 'source.pdf')
    digest = sha256(source.read_bytes()).digest()
    control = JobControl()
    def translate(request, cancel):
        control.cancel()
        return SimpleNamespace(translated_text='Перевод')
    with pytest.raises(TranslationCancelledError):
        pdf_job([source], control=control, translate=translate).run()
    def fail(*args):
        raise PdfError(PdfKind.CORRUPTED_PDF)
    monkeypatch.setattr(PdfDocument, 'validate', fail)
    with pytest.raises(PdfError):
        pdf_job([source]).run()
    assert list(tmp_path.iterdir()) == [source]
    assert sha256(source.read_bytes()).digest() == digest


def test_limits_and_font_coverage(tmp_path):
    source = make_pdf(tmp_path / 'source.pdf', pages=2)
    for limits in [replace(PdfLimits(), max_pages=1), replace(PdfLimits(), max_file_bytes=100),
                   replace(PdfLimits(), max_objects_per_page=1), replace(PdfLimits(), max_text_chars=2)]:
        with pytest.raises(PdfError) as error:
            PdfDocument(source, limits)
        assert error.value.kind == PdfKind.UNSUPPORTED_PDF
    resolver = FontResolver()
    try:
        assert resolver.resolve('Latin Кириллица 中文').covers('Latin Кириллица 中文')
    finally:
        resolver.close()
