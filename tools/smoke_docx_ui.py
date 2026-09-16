"""Real Windows Qt DOCX smoke with synthetic inputs, local models and isolated settings."""
import json
import os
from hashlib import sha256
from pathlib import Path
import sys
from time import monotonic

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from PySide6.QtCore import QSettings, QTimer, Qt, QMimeData, QUrl, QPointF
from PySide6.QtGui import QDropEvent, QDragEnterEvent
from PySide6.QtWidgets import QApplication

from app.documents.docx_document import validate_docx
from app.gui.main_window import MainWindow
from app.gui.styles.theme import load_stylesheet
from app.models.translation_job import JobState


def main():
    root = Path('build/aw05-gui-smoke').resolve()
    root.mkdir(parents=True, exist_ok=True)
    report_dir = Path('docs/qa')
    report_dir.mkdir(parents=True, exist_ok=True)
    app = QApplication([])
    app.setOrganizationName('CHOPIK Team QA')
    app.setApplicationName('TreeTranslate AW05 DOCX QA')
    QSettings.setDefaultFormat(QSettings.Format.IniFormat)
    QSettings.setPath(QSettings.Format.IniFormat, QSettings.Scope.UserScope, str(root / 'settings'))
    app.setStyleSheet(load_stylesheet())
    window = MainWindow()
    window.setAttribute(Qt.WidgetAttribute.WA_DontShowOnScreen, True)
    window.settings.save_value('general/output_location', 'custom')
    window.settings.save_value('general/output_path', str(root / 'outputs'))
    window.settings.save_value('general/output_template', '{name}_{lang}')
    window.preferences.set_languages('Определить автоматически', 'Русский')
    window.preferences.set_profile('Баланс')
    window.preferences.set_translate_directories(False)
    window.show()
    engine = window.translation_service.engine
    scenarios = [('en_cpu', 'CPU', 'Save the configuration file before restarting the application.'),
                 ('zh_auto', 'Auto', '请在重新启动应用程序之前保存配置文件。')]
    if engine.devices.gpu_available():
        scenarios.append(('zh_gpu', 'GPU', '请在重新启动应用程序之前保存配置文件。'))
    scenarios.extend([('folder', 'Auto', 'Save the configuration file.'),
                      ('pause', 'CPU', 'Save the configuration file.'),
                      ('cancel', 'CPU', 'Save the configuration file.')])
    reports, errors, calls = [], [], []
    original_translate = engine.translate
    def translate(request, cancelled):
        result = original_translate(request, cancelled)
        calls.append({'source': result.source_language, 'target': result.target_language,
                      'backend': result.backend, 'device': result.device, 'ms': result.duration_ms,
                      'fallback': result.fallback_used})
        return result
    engine.translate = translate
    current = {}

    def fail(message):
        errors.append(message)
        QTimer.singleShot(0, window.close)

    def begin():
        if not scenarios:
            window.close()
            return
        name, device, text = scenarios.pop(0)
        path = root / (name + '.docx')
        doc = Document()
        doc.add_heading('User guide' if name.startswith('en') else text, 1)
        paragraph = doc.add_paragraph()
        paragraph.add_run(text[:len(text)//2]).bold = True
        paragraph.add_run(text[len(text)//2:]).italic = True
        doc.add_table(rows=1, cols=1).cell(0, 0).text = text
        doc.sections[0].header.paragraphs[0].text = text
        doc.sections[0].footer.paragraphs[0].text = text
        if name in {'pause', 'cancel'}:
            for _ in range(8):
                doc.add_paragraph(text)
        doc.save(path)
        inputs = [path]
        sources = [path]
        if name == 'folder':
            folder = root / 'User guide'
            child = folder / 'Installation'
            child.mkdir(parents=True, exist_ok=True)
            first, second = folder / 'first.docx', child / 'second.docx'
            doc.save(first)
            doc.save(second)
            inputs, sources = [folder], [first, second]
        current.clear()
        current.update(name=name, device=device, path=path, hash=sha256(path.read_bytes()).hexdigest(),
                       source_hashes={str(p): sha256(p.read_bytes()).hexdigest() for p in sources},
                       call_start=len(calls), started=monotonic(), paused=False)
        window.preferences.set_device(device)
        window.preferences.set_translate_directories(name == 'folder')
        mime = QMimeData()
        mime.setUrls([QUrl.fromLocalFile(str(p)) for p in inputs])
        zone = window.file_page.drop_zone
        drag = QDragEnterEvent(zone.rect().center(), Qt.DropAction.CopyAction, mime,
                               Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
        QApplication.sendEvent(zone, drag)
        drop = QDropEvent(QPointF(zone.rect().center()), Qt.DropAction.CopyAction, mime,
                          Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
        QApplication.sendEvent(zone, drop)
        if not drop.isAccepted():
            fail('DOCX drag/drop rejected')

    def state_changed(state):
        try:
            if state == JobState.READY:
                QTimer.singleShot(0, window.file_page.start_button.click)
            elif state in {JobState.COMPLETED, JobState.CANCELLED}:
                expected = JobState.CANCELLED if current['name'] == 'cancel' else JobState.COMPLETED
                assert state == expected, (state, expected)
                after = sha256(current['path'].read_bytes()).hexdigest()
                assert after == current['hash']
                source_after = {p: sha256(Path(p).read_bytes()).hexdigest() for p in current['source_hashes']}
                assert source_after == current['source_hashes']
                outputs = window.file_page.progress.output_paths
                if state == JobState.COMPLETED:
                    assert outputs
                    validate_docx(outputs[0])
                    assert any('а' <= c.lower() <= 'я' for c in Document(outputs[0]).paragraphs[1].text)
                    assert window.file_page.progress.bar.value() == 100
                    assert window.file_page.progress.open_file.isEnabled()
                    assert window.file_page.progress.show_output.isEnabled()
                    if current['name'] == 'folder':
                        assert len(outputs) == 2
                        assert outputs[1].parent.parent == outputs[0].parent
                        assert outputs[1].parent.name != 'Installation'
                        validate_docx(outputs[1])
                else:
                    assert not outputs
                    assert len(calls) > current['call_start'], 'Cancel must exercise an active translation'
                    assert not list((root / 'outputs').glob('.treetranslate-*.docx'))
                if current['name'] == 'pause':
                    assert current['paused']
                record = {'scenario': current['name'], 'requested_device': current['device'],
                          'state': state.name, 'source_sha256_before': current['hash'], 'source_sha256_after': after,
                          'source_hashes_before': current['source_hashes'], 'source_hashes_after': source_after,
                          'outputs': [str(p) for p in outputs], 'elapsed_seconds': round(monotonic()-current['started'], 3),
                          'calls': calls[current['call_start']:]}
                reports.append(record)
                print(json.dumps(record, ensure_ascii=True), flush=True)
                def capture_and_continue():
                    window.grab().save(str(report_dir / f"aw05-{current['name']}.png"))
                    begin()
                QTimer.singleShot(50, capture_and_continue)
        except Exception as error:
            fail(repr(error))

    def progress_changed(progress):
        if current.get('name') == 'cancel' and progress.processed == 1:
            QTimer.singleShot(5, window.file_page.progress.cancel.click)
        if current.get('name') == 'pause' and progress.processed == 1 and not current['paused']:
            current['paused'] = True
            window.file_page.progress.pause.click()
            QTimer.singleShot(150, verify_pause)

    def verify_pause():
        try:
            assert window.translation_service.state == JobState.PAUSED
            previous = len(calls)
            def resume():
                try:
                    # One bounded inference already in flight may finish while paused.
                    assert len(calls) <= previous + 1
                    assert engine.runtime._pins == 1
                    window.file_page.progress.pause.click()
                except Exception as error:
                    fail(repr(error))
            QTimer.singleShot(250, resume)
        except Exception as error:
            fail(repr(error))

    window.translation_service.state_changed.connect(state_changed)
    window.translation_service.progress_changed.connect(progress_changed)
    window.translation_service.files.failed.connect(fail)
    timeout = QTimer(singleShot=True, interval=120000)
    timeout.timeout.connect(lambda: fail('DOCX GUI smoke timed out'))
    timeout.start()
    QTimer.singleShot(0, begin)
    app.exec()
    timeout.stop()
    (report_dir / 'aw05-docx-gui.json').write_text(json.dumps({'scenarios': reports, 'errors': errors}, indent=2), encoding='utf-8')
    assert not errors, errors
    assert len(reports) >= 4
    print('Real DOCX Qt drag/drop, CPU/Auto/GPU, pause/resume, cancel, validation and hashes passed', flush=True)


if __name__ == '__main__':
    main()
