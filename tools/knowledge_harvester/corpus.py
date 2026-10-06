"""Offline document intake using the existing Harvester store and provenance.

Candidate text is a retrieval hypothesis, never a bilingual assertion. Sources
are read-only; review records bind acceptance to the evidence digest.
"""
from datetime import datetime, timezone
from io import BytesIO
import json
from pathlib import Path
import re
from zipfile import ZipFile

from .models import HarvestError
from .normalization import normalize
from .provenance import canonical_json, digest_json, file_hash

VERSION = '0.8.8'
ORIGINS = {'AUTHORED', 'USER_PROVIDED', 'EXISTING_PROJECT_FIXTURE',
           'LICENSED_OPEN_SOURCE', 'MODEL_SUGGESTED', 'DERIVED_FROM_EXISTING_KNOWLEDGE'}
STATES = {'CANDIDATE', 'REVIEWED', 'VERIFIED', 'REJECTED', 'AMBIGUOUS', 'DEPRECATED'}
TYPES = {'TERM', 'COMPOUND', 'PHRASE', 'TEMPLATE', 'FULL_SEGMENT', 'ALIAS', 'ABBREVIATION', 'CONCEPT_RELATION'}


def _schema(con):
    # Additive dev tables; the original store schema and import/export stay intact.
    con.executescript('''
      CREATE TABLE IF NOT EXISTS corpus_documents(id TEXT PRIMARY KEY,payload TEXT NOT NULL);
      CREATE TABLE IF NOT EXISTS document_evidence(candidate TEXT,document TEXT,offset INTEGER,
        PRIMARY KEY(candidate,document,offset));
      CREATE INDEX IF NOT EXISTS corpus_source_lookup ON document_evidence(document);
    ''')


def texts(data, suffix):
    if suffix == '.pdf':
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(data) as doc:
            if len(doc) > 2000:raise HarvestError('page_budget')
            out=[]
            for page in doc:
                text=page.get_textpage()
                try:out.extend(text.get_text_range().splitlines())
                finally:text.close();page.close()
            return out
    if suffix == '.docx':
        from docx import Document
        from app.documents.zip_archive import validate_members
        with ZipFile(BytesIO(data)) as archive:validate_members(archive.infolist())
        doc=Document(BytesIO(data));out=[p.text for p in doc.paragraphs]
        out.extend(c.text for t in doc.tables for r in t.rows for c in r.cells)
        return out
    if suffix in {'.txt', '.tsv'}:return data.decode('utf-8-sig').splitlines()
    raise HarvestError('unsupported_corpus_type')


def candidate(store, *, source, target='', domain='general', subdomains=(), kind='TERM',
              origin='USER_PROVIDED', provenance=None, concept_id='', aliases=(),source_language='zh',target_language='ru'):
    if origin not in ORIGINS or kind not in TYPES:raise HarvestError('candidate_metadata')
    original_source=source;source=normalize(source);target=normalize(target)
    if not source or len(source)>256 or len(target)>512:raise HarvestError('candidate_length')
    scope=sorted(set(subdomains));uid=digest_json([source,target,domain,scope,kind,source_language,target_language])
    evidence=dict(provenance or {},origin=origin,harvester_version=VERSION)
    with store.connect() as con:
        _schema(con)
        # A declared alias is not a fresh concept. Reuse only an identical
        # bilingual assertion/context; conflicting aliases remain review items.
        for row in con.execute('SELECT c.payload FROM candidate_aliases a JOIN candidates c ON c.id=a.candidate WHERE a.alias=? AND a.domain=?',(source,domain)):
            previous=json.loads(row[0])
            if (previous['ru'],previous.get('subdomains',[]),previous['type'],previous.get('source_language','zh'),previous.get('target_language','ru'))==(target,scope,kind,source_language,target_language):
                uid=previous['id'];break
        existing=con.execute('SELECT payload FROM candidates WHERE id=?',(uid,)).fetchone()
        payload=json.loads(existing[0]) if existing else dict(id=uid,zh=source,ru=target,domain=domain,
            subdomains=scope,type=kind,concept_id=concept_id or 'corpus:'+digest_json([source,domain,scope])[:20],
            status='CANDIDATE',provenance=[],reasons=[],manual_review=False,aliases=list(aliases),original_source=original_source,source_language=source_language,target_language=target_language)
        if evidence not in payload['provenance']:
            payload['provenance'].append(evidence)
            if existing and payload['status']=='VERIFIED':payload['status']='REVIEWED';payload['manual_review']=False
        # Different contexts preserve polysemy; only overlapping contexts conflict.
        for other in con.execute('SELECT id,payload FROM candidates WHERE id IN (SELECT candidate FROM candidate_aliases WHERE alias=?) AND domain=? AND ru!=? AND ru!=?',
                                 (source,domain,target,'')):
            row=json.loads(other['payload'])
            if (row.get('source_language','zh'),row.get('target_language','ru'))!=(source_language,target_language):continue
            overlap=not scope or not row.get('subdomains') or bool(set(scope)&set(row['subdomains']))
            if target and overlap:
                payload['status']='AMBIGUOUS'
                con.execute('INSERT OR IGNORE INTO conflicts VALUES(?,?,?)',(uid,other['id'],'context_target_conflict'))
        con.execute('INSERT OR REPLACE INTO candidates VALUES(?,?,?,?,?,?,?,?)',
                    (uid,source,target,domain,payload['status'],'DOCUMENT_CANDIDATE',0,canonical_json(payload)))
        for alias in (source,*aliases):
            con.execute('INSERT OR IGNORE INTO candidate_aliases VALUES(?,?,?)',(normalize(alias),domain,uid))
    return uid


def review(store, uid, state, *, reviewer, reason, evidence_sha256):
    if state not in STATES or not reviewer or not reason:raise HarvestError('review_metadata')
    with store.connect() as con:
        row=con.execute('SELECT payload FROM candidates WHERE id=?',(uid,)).fetchone()
        if not row:raise HarvestError('unknown_candidate')
        payload=json.loads(row[0])
        if digest_json(payload['provenance'])!=evidence_sha256:raise HarvestError('changed_evidence')
        if state=='VERIFIED':
            if payload['status']!='REVIEWED' or not payload['ru']:raise HarvestError('review_required')
            if any(p['origin']=='MODEL_SUGGESTED' for p in payload['provenance']):raise HarvestError('model_requires_independent_source')
            if con.execute('SELECT 1 FROM conflicts WHERE candidate=? OR other=?',(uid,uid)).fetchone():raise HarvestError('unresolved_conflict')
        payload.update(status=state,manual_review=True,review=dict(reviewer=reviewer,reason=reason,evidence_sha256=evidence_sha256))
        con.execute('UPDATE candidates SET status=?,payload=? WHERE id=?',(state,canonical_json(payload),uid))
        con.execute('INSERT OR REPLACE INTO reviews VALUES(?,?,?)',(uid,state,canonical_json(payload['review'])))


def intake(store, path, *, corpus_id, language, origin, permission='REVIEW_REQUIRED', development_manifest=None):
    from app.documents.zip_archive import inventory
    from app.documents.control import JobControl
    from app.knowledge.profile import DocumentProfiler
    from app.knowledge.segments import SegmentClassifier
    if origin not in ORIGINS or language=='auto' or not corpus_id:raise HarvestError('intake_metadata')
    path=Path(path).resolve();docs=[]
    manifest=json.loads(Path(development_manifest).read_text('utf-8')) if development_manifest else None
    allowed={d['member']:d['sha256'] for d in manifest['documents']} if manifest else None
    files=sorted(p for p in path.rglob('*') if p.is_file() and not p.is_symlink()) if path.is_dir() else [path]
    if len(files)>100000:raise HarvestError('corpus_file_budget')
    profiler=DocumentProfiler()
    for file in files:
        digest=file_hash(file)
        if file.suffix.lower()=='.zip':
            if manifest and digest!=manifest['zip_sha256']:raise HarvestError('changed_corpus_container')
            members=inventory(file,JobControl())
            def archive_inputs():
                from tempfile import TemporaryDirectory
                from app.documents.zip_archive import extract
                with ZipFile(file) as archive:
                    for member in members:
                        if allowed is not None and member.name not in allowed:continue
                        if not member.directory and Path(member.name).suffix.lower() in {'.pdf','.docx'}:
                            # Bounded temporary member rather than archive.read.
                            with TemporaryDirectory(prefix='TreeTranslate-corpus-') as root:
                                extract(file,Path(root),(member,),JobControl(),archive=archive)
                                working=Path(root)/member.name
                                if allowed is not None and file_hash(working)!=allowed[member.name]:raise HarvestError('changed_corpus_member')
                                if member.size>256*1024*1024:raise HarvestError('corpus_member_parse_budget')
                                yield member.name,working.read_bytes()
            inputs=archive_inputs()
        elif file.suffix.lower() in {'.pdf','.docx','.txt','.tsv'}:
            if file.stat().st_size>256*1024*1024:raise HarvestError('corpus_member_parse_budget')
            inputs=[('',file.read_bytes())]
        else:continue
        for member,data in inputs:
            from hashlib import sha256
            suffix=Path(member or file.name).suffix.lower();segments=texts(data,suffix)
            if sum(map(len,segments))>8_000_000:raise HarvestError('corpus_text_budget')
            profile=profiler.profile(language,'ru' if language!='ru' else 'en',segments=segments,filename=member or file.name)
            uid=digest_json([corpus_id,str(file),digest,member,sha256(data).hexdigest()])
            payload=dict(id=uid,corpus_id=corpus_id,source_path=str(file),archive_member=member,
                language=language,detected_domain=profile.primary_domain,subdomains=list(dict(profile.subdomains)),
                document_type=suffix.lstrip('.'),sha256=sha256(data).hexdigest(),container_sha256=digest,
                source_type=origin,usage_permission_status=permission,processed_timestamp=datetime.now(timezone.utc).isoformat(),
                harvester_version=VERSION,segments=len(segments),ocr_used=False)
            with store.connect() as con:
                _schema(con);con.execute('INSERT OR IGNORE INTO corpus_documents VALUES(?,?)',(uid,canonical_json(payload)))
            for offset,text in enumerate(segments):
                text=normalize(text)
                if not text or len(text)>96 or not re.search(r'[\w\u4e00-\u9fff]',text):continue
                kind=SegmentClassifier.classify(text).value
                ctype='ABBREVIATION' if re.fullmatch(r'[A-Z]{2,8}',text) else 'PHRASE' if kind in {'PROCEDURE_STEP','WARNING','CONDITION'} else 'COMPOUND' if len(text)>4 else 'TERM'
                cid=candidate(store,source=text,domain=profile.primary_domain,subdomains=payload['subdomains'],kind=ctype,
                    origin=origin,source_language=language,target_language='ru' if language!='ru' else 'en',
                    provenance=dict(document_id=uid,sha256=payload['sha256'],offset=offset,permission=permission))
                with store.connect() as con:con.execute('INSERT OR IGNORE INTO document_evidence VALUES(?,?,?)',(cid,uid,offset))
                action=re.fullmatch(r'(?:\d+[.、]\s*)?(拆下|拆卸|安装|更换|检查|拧紧|测量|断开|连接)([\u4e00-\u9fff]{2,24})[。.]?',text)
                if action:
                    evidence=dict(document_id=uid,sha256=payload['sha256'],offset=offset,permission=permission)
                    for term,typ in [(action[2],'COMPOUND'),(action[1]+'{X}','TEMPLATE')]:
                        candidate(store,source=term,domain=profile.primary_domain,subdomains=payload['subdomains'],kind=typ,origin=origin,
                            source_language=language,target_language='ru' if language!='ru' else 'en',provenance=evidence)
            docs.append(payload)
        if file_hash(file)!=digest:raise HarvestError('source_changed')
    return docs


def harvest_manifest(store, cache_path, manifest_path):
    """Harvest native developer evidence through a frozen development allowlist.

    Conversion cache has bounded PDFs' native lines, not translations. The
    cache digest is frozen with the split, and holdout paths never enter Store.
    Frequency and priority explain review order; neither asserts correctness.
    """
    from collections import Counter, defaultdict
    import sqlite3
    from app.knowledge.segments import SegmentClassifier
    manifest=json.loads(Path(manifest_path).read_text('utf-8'))
    manifest_digest=file_hash(manifest_path)
    if file_hash(cache_path)!=manifest['native_cache_sha256']:raise HarvestError('changed_native_cache')
    allowed={d['member']:d for d in manifest['documents']}
    records={};rejected=Counter();documents=[];evidence=[]
    with sqlite3.connect(cache_path) as cache:
        for member,digest,size,raw in cache.execute('SELECT * FROM documents ORDER BY member'):
            if member not in allowed:continue
            item=allowed[member]
            if digest!=item['sha256']:raise HarvestError('changed_corpus_member')
            document=json.loads(raw);lines=document.get('lines',[])
            uid=digest_json([manifest['corpus_id'],member,digest])
            payload=dict(id=uid,corpus_id=manifest['corpus_id'],archive_member=member,sha256=digest,
                container_sha256=manifest['source_sha256'],language='zh',detected_domain=item['domain'],
                subdomains=item['subdomains'],source_type='USER_PROVIDED',usage_permission_status='CANDIDATES_ONLY',
                segments=len(lines),ocr_used=False,harvester_version=VERSION)
            documents.append(payload)
            sources=[(i,line,'content') for i,line in enumerate(lines)]
            sources += [(-i-1,Path(part).stem,'folder') for i,part in enumerate(Path(member).parts[2:])]
            for offset,line,origin in sources:
                line=normalize(line);typ=SegmentClassifier.classify(line).value
                runs=re.findall(r'[\u4e00-\u9fff]{2,32}',line)
                if not runs:rejected['no_Chinese_or_protected_ID']+=1;continue
                options=[(run,'TERM' if len(run)<=3 else 'COMPOUND') for run in runs]
                if len(line)<=96 and typ in {'PROCEDURE_STEP','WARNING','CONDITION'}:
                    options.append((line,'PHRASE'))
                action=re.match(r'^(?:\d+[.、]\s*)?(拆下|拆卸|安装|更换|检查|拧紧|测量|断开|连接|分离|调整|清洁)([\u4e00-\u9fff]{2,24})(?:[（(。.]|$)',line)
                if action:options.extend([(action[2],'COMPOUND'),(action[1]+'{X}','TEMPLATE')])
                for source,kind in options:
                    if len(source)>96:rejected['length']+=1;continue
                    key=(source,kind)
                    record=records.setdefault(key,dict(source=source,type=kind,total_frequency=0,
                        documents=set(),subdomains=Counter(),segment_types=Counter(),origins=Counter(),samples=[]))
                    record['total_frequency']+=1;record['documents'].add(uid)
                    record['subdomains'].update([item['stratum']]);record['segment_types'].update([typ]);record['origins'].update([origin])
                    if len(record['samples'])<6 and not any(s['document_id']==uid for s in record['samples']):
                        record['samples'].append(dict(document_id=uid,member=member,sha256=digest,offset=offset))
    with store.connect() as con:
        _schema(con)
        con.executemany('INSERT OR REPLACE INTO corpus_documents VALUES(?,?)',[(d['id'],canonical_json(d)) for d in documents])
    output=[]
    for (source,kind),record in records.items():
        branches=[b for b,n in record['subdomains'].most_common() if b!='general' and n>=max(1,record['total_frequency']*.2)]
        domains=Counter(b.split('.')[0] for b in branches)
        domain=domains.most_common(1)[0][0] if domains else 'general'
        n=len(record['documents'])
        factors=dict(document_frequency=min(20,n)*2,compound_specificity=min(10,len(source)),
            label_evidence=5 if any(t in record['segment_types'] for t in ('HEADING','DIAGRAM_LABEL','COMPONENT_LABEL','TABLE_CELL')) else 0,
            action_structure=6 if kind=='TEMPLATE' else 0,ambiguous_short=-10 if len(source)==2 else 0,
            single_occurrence=-5 if n==1 else 0,context_conflict=-5 if len(domains)>1 else 0)
        provenance=dict(corpus_id=manifest['corpus_id'],development_manifest_sha256=manifest_digest,
            permission='CANDIDATES_ONLY',evidence_count=n,total_frequency=record['total_frequency'],samples=record['samples'])
        cid=candidate(store,source=source,kind=kind,domain=domain,subdomains=branches,
            origin='USER_PROVIDED',provenance=provenance)
        with store.connect() as con:
            con.executemany('INSERT OR IGNORE INTO document_evidence VALUES(?,?,?)',[(cid,d,0) for d in record['documents']])
        output.append(dict(id=cid,source=source,type=kind,document_frequency=n,total_frequency=record['total_frequency'],
            subdomain_frequency=dict(record['subdomains']),segment_types=dict(record['segment_types']),
            contexts=dict(record['origins']),samples=record['samples'],priority_factors=factors,
            review_priority=sum(factors.values()),status='CANDIDATE'))
    if file_hash(cache_path)!=manifest['native_cache_sha256']:raise HarvestError('changed_native_cache')
    return dict(documents=len(documents),candidates=sorted(output,key=lambda r:(-r['review_priority'],-r['total_frequency'],r['source'])),
        rejections=dict(rejected),holdout_excluded=True,automatic_verified=0)


def export(store, destination):
    with store.connect() as con:
        rows=[json.loads(r[0]) for r in con.execute('SELECT payload FROM candidates ORDER BY id')]
    Path(destination).write_text('\n'.join(canonical_json(r) for r in rows)+'\n','utf-8')
    return len(rows)


def reviewed_entries(store):
    """Adapter to the existing pack writer, with redistribution evidence gate."""
    with store.connect() as con:
        for row in con.execute("SELECT payload FROM candidates WHERE status='VERIFIED' ORDER BY id"):
            payload=json.loads(row[0]);review_record=payload.get('review',{})
            if not payload.get('manual_review') or digest_json(payload['provenance'])!=review_record.get('evidence_sha256'):continue
            if payload['type'] not in {'TERM','COMPOUND','PHRASE','FULL_SEGMENT','ABBREVIATION'}:continue
            if payload['type']=='FULL_SEGMENT' and len(payload['zh'])>48:continue
            if any(p['origin']=='MODEL_SUGGESTED' or p.get('permission') not in {'AUTHORED_FOR_PROJECT','APPROVED_FOR_REDISTRIBUTION'} for p in payload['provenance']):continue
            yield dict(source_term=payload['zh'],target_term=payload['ru'],source_language=payload.get('source_language','zh'),
                target_language=payload.get('target_language','ru'),domain=payload['domain'],status='BUILTIN',
                variants=payload.get('aliases',[]),notes=canonical_json(dict(type=payload['type'].lower(),concept_id=payload['concept_id'],
                    subdomains=payload['subdomains'],review_status='VERIFIED',review_candidate=payload['id'])),
                provenance=canonical_json(payload['provenance'])[:1024])
