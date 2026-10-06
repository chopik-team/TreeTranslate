"""Small real offline document baseline, including OCR worker RAM and stage times."""
import argparse
from dataclasses import asdict
from hashlib import sha256
import json
import os
from pathlib import Path
import socket
import sys
from time import perf_counter

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--phase',choices=['baseline','final'],required=True);args=parser.parse_args()
    build=ROOT/'build/aw08';qa=ROOT/'docs/qa/aw08';build.mkdir(parents=True,exist_ok=True)
    os.environ['TREETRANSLATE_TM_PATH']=str(build/'qa-tm.db');os.environ['TREETRANSLATE_GLOSSARY_PATH']=str(build/'qa-glossary.db')
    from docx import Document
    from app.documents.control import JobControl
    from app.documents.job import DocumentConfig,DocumentJob
    from app.documents.scanner import scan_sources
    import app.documents.job as jobs
    from app.engine.factory import create_translation_engine
    from app.engine.types import DevicePreference,PerformanceProfile
    from app.ocr.router.ocr_router import OcrRouter
    from tools.pdf_fixtures import make_pdf
    from tools.ocr_fixtures import rasterize,make_mixed
    from tools.benchmark_translation import MemorySampler,gpu_used_mb
    import psutil
    import pypdfium2 as pdfium
    sentence='Do not start the pump before the valve is open.'
    doc=Document();doc.add_heading('Installation guide',0);doc.add_paragraph(sentence)
    doc.add_paragraph('Record 123 and 42 mm.',style='List Bullet');doc.add_paragraph('')
    doc.add_table(rows=1,cols=2).rows[0].cells[0].text='Cooling water pressure'
    doc.tables[0].cell(0,1).text='12 bar'
    doc.sections[0].header.paragraphs[0].text='Safety instructions'
    doc.sections[0].footer.paragraphs[0].text='Save the configuration file.'
    doc.save(build/'quality.docx')
    native=make_pdf(build/'native.pdf',sentence,lines=3,background=False)
    scans={'scan':rasterize(native,build/'scan.pdf'), 'rotated':rasterize(native,build/'rotated.pdf',angle=90),
           'poor':rasterize(native,build/'poor.pdf',dpi=110), 'mixed':make_mixed(build/'mixed.pdf')}
    paths={'docx':build/'quality.docx','native':native,**scans}
    attempts=[]
    def blocked(*a,**k):attempts.append(True);raise RuntimeError('QA blocked network')
    socket.getaddrinfo=blocked;socket.create_connection=blocked
    class TreeMemory(MemorySampler):
        def rss(self):
            total=0
            for process in [self.process,*self.process.children(recursive=True)]:
                try:total+=process.memory_info().rss
                except psutil.Error:pass
            return total/1024**2
    engine=create_translation_engine();rows=[];captured=[];ocr=[];write_seconds=[]
    original_open=jobs.open_document;original_recognize=OcrRouter.recognize
    def opened(*a,**k):
        tick=perf_counter();document=original_open(*a,**k);captured.append((document,perf_counter()-tick))
        original_write=document.write
        def write(*a,**k):
            tick=perf_counter()
            try:return original_write(*a,**k)
            finally:write_seconds.append(perf_counter()-tick)
        document.write=write
        return document
    def recognize(self,request):
        result=original_recognize(self,request)
        ocr.append({'duration':result.duration,'timings':result.timings,'backend':result.backend,'device':result.device,'segments':len(result.segments)})
        return result
    jobs.open_document=opened;OcrRouter.recognize=recognize
    matrix=[('docx','auto',3),('native','auto',3),('scan','cpu',2),('scan','gpu',2),('mixed','auto',1),('rotated','auto',1),('poor','auto',1)]
    try:
        for name,device,repeats in matrix:
            for iteration in range(repeats):
                captured.clear();ocr.clear();write_seconds.clear();progress=[];warnings=[];translation_ms=[]
                path=paths[name];digest=sha256(path.read_bytes()).hexdigest();control=JobControl()
                memory=TreeMemory();memory.start();cpu_before=psutil.Process().cpu_times();start=perf_counter()
                def translate(request,cancel):
                    tick=perf_counter()
                    try:return engine.translate(request,cancel)
                    finally:translation_ms.append((perf_counter()-tick)*1000)
                row={'case':name,'requested_device':device,'iteration':iteration,'model_state':'cold' if not rows else 'mixed retained/released; see stages', 'rss_before_mb':memory.rss()}
                try:
                    with engine.runtime.keep_warm():
                        output,=DocumentJob(scan_sources([path],control).files,
                            DocumentConfig(source='en',target='ru',device=DevicePreference(device),profile=PerformanceProfile.BALANCED,
                                           output=build/f'outputs-{args.phase}'),control,translate,engine.languages.resolve,
                            progress.append,warning=warnings.append,before_ocr=engine.runtime.release_models).run()
                    row.update(output=str(output),source_unchanged=digest==sha256(path.read_bytes()).hexdigest(),
                               completed=True,progress_events=len(progress),last_percent=progress[-1].percent,warnings=list(dict.fromkeys(warnings)))
                    document=captured[-1][0]
                    row['transcript']=[{'source':s.text,'output':s.translated} for s in document.segments]
                    if name!='docx':
                        with pdfium.PdfDocument(output) as pdf:row['pages']=len(pdf)
                        document.validate(output)
                    else:
                        result=Document(output);row['structure']={'tables':len(result.tables),'paragraphs':len(result.paragraphs),'header':result.sections[0].header.paragraphs[0].text,'footer':result.sections[0].footer.paragraphs[0].text}
                except Exception as error:row.update(completed=False,error_type=type(error).__name__,error=str(error))
                row.update(seconds=perf_counter()-start,translation_seconds=sum(translation_ms)/1000,
                           open_extract_seconds=sum(x[1] for x in captured),write_seconds=sum(write_seconds),ocr=list(ocr),
                           rss_after_mb=memory.rss(),gpu_total_mb=gpu_used_mb())
                cpu_after=psutil.Process().cpu_times()
                row['parent_cpu_seconds']=cpu_after.user+cpu_after.system-cpu_before.user-cpu_before.system
                memory.close();row['tree_peak_rss_mb']=memory.peak
                rows.append(row)
                (qa/f'documents-{args.phase}.json').write_text(json.dumps({'network_attempts_parent':len(attempts),'ocr_worker_guard':'existing offline worker guard','results':rows},ensure_ascii=False,indent=2),encoding='utf-8')
                print(name,device,iteration,row['completed'],round(row['seconds'],3),flush=True)
    finally:
        engine.shutdown();jobs.open_document=original_open;OcrRouter.recognize=original_recognize


if __name__=='__main__':main()
